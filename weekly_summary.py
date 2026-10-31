"""
weekly_summary.py - emails YOU a weekly report: what was found, what you applied to, and which
sources actually led to replies. Reads only your own files. Sends at most once every 7 days.

For the replies column to work, write short notes in the jobs Excel notes column, for example:
  "Applied 20 Oct"   then later   "Applied 20 Oct, interview 28 Oct"   or   "Applied 20 Oct, rejected"
Words counted as a positive reply: interview, replied, reply, call, shortlist, referred, referral, offer,
assessment, screening.  Negative: reject, declined, not selected, no reply, no response.
Run: python weekly_summary.py   (add --print to preview, --force to send now)
"""
import os
import re
import smtplib
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from email.mime.text import MIMEText

from openpyxl import load_workbook

from followup_reminder import (EMAIL_ADDRESS, EMAIL_PASSWORD, KIT_PATH, RECIPIENT, SMTP_PORT,  # noqa: E402
                               SMTP_SERVER, date_from_note, due_from_applications, due_from_kit)
from networking_helper import parse_sent_date  # noqa: E402

DB_PATH = os.getenv("DB_PATH", "jobs.db")
MIN_SCORE = int(os.getenv("MIN_EMAIL_MATCH_SCORE", "60"))
GUARD = "weekly_summary_last_sent.txt"
NEGATIVE = re.compile(r"reject|declin|not selected|no reply|no response", re.I)
POSITIVE = re.compile(r"interview|replied|reply|call|shortlist|referred|referral|offer|assessment|screening", re.I)


def source_of(row_site, row_channel):
    return (row_site or "").strip() or (row_channel or "").strip() or "unknown"


def collect(today):
    week_start = today - timedelta(days=7)
    data = {"found": 0, "matches": 0, "by_source_matches": Counter(), "applied_week": 0,
            "apps": defaultdict(lambda: [0, 0, 0]), "kit_status": Counter(), "kit_week": 0}
    if os.path.exists(DB_PATH):
        conn = sqlite3.connect(DB_PATH)
        since = week_start.isoformat()
        data["found"] = conn.execute("SELECT COUNT(*) FROM jobs WHERE created_at >= ?", (since,)).fetchone()[0]
        for site, chan, n in conn.execute(
                "SELECT source_site, source_channel, COUNT(*) FROM jobs WHERE created_at >= ? AND scored=1 "
                "AND match_score >= ? GROUP BY source_site, source_channel", (since, MIN_SCORE)):
            data["by_source_matches"][source_of(site, chan)] += n
        data["matches"] = sum(data["by_source_matches"].values())
        for site, chan, note in conn.execute(
                "SELECT source_site, source_channel, user_notes FROM jobs WHERE user_notes LIKE '%pplied%'"):
            note = note or ""
            d = date_from_note(note)
            if d and d >= week_start:
                data["applied_week"] += 1
            row = data["apps"][source_of(site, chan)]
            row[0] += 1
            if NEGATIVE.search(note):
                row[2] += 1
            elif POSITIVE.search(note):
                row[1] += 1
        conn.close()
    if os.path.exists(KIT_PATH):
        try:
            wb = load_workbook(KIT_PATH, data_only=True)
            for sheet in ("Outreach_Drafts", "Target_Companies"):
                ws = wb[sheet]
                hdr = [c.value for c in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    st = str(row[hdr.index("Status")] or "").strip()
                    if st and st.lower() != "not sent":
                        data["kit_status"][st] += 1
                    d = parse_sent_date(row[hdr.index("Date sent")])
                    if d and d >= week_start:
                        data["kit_week"] += 1
        except Exception as exc:
            print(f"[weekly] cannot read {KIT_PATH}: {exc}")
    data["overdue"] = len(due_from_kit(today)) + len(due_from_applications(today))
    return data


def build_body(d, today):
    L = [f"Weekly job search summary, week ending {today:%d %b %Y}", "",
         f"Jobs found this week: {d['found']}   |   Strong matches (score {MIN_SCORE}+): {d['matches']}",
         f"Applications you logged this week: {d['applied_week']}",
         f"Networking messages you logged this week: {d['kit_week']}",
         f"Follow-ups overdue now: {d['overdue']}", ""]
    L.append("Where this week's strong matches came from:")
    L += [f"  {s}: {n}" for s, n in d["by_source_matches"].most_common()] or ["  none"]
    L += ["", "Applications by source, all time (applied / positive reply / negative):"]
    if d["apps"]:
        for s, (a, p, n) in sorted(d["apps"].items(), key=lambda x: -x[1][0]):
            L.append(f"  {s}: {a} applied, {p} positive, {n} negative")
    else:
        L.append('  none logged yet. Write "Applied 20 Oct" in the notes column when you apply.')
    L += ["", "Networking statuses: " + (", ".join(f"{k}: {v}" for k, v in d["kit_status"].items()) or "none yet"), ""]
    tips = []
    if d["overdue"]:
        tips.append(f"Send the {d['overdue']} overdue follow-up(s) first. They are the cheapest wins.")
    if d["applied_week"] < 5:
        tips.append("Apply to more of the Top_60_Plus roles; the aim is a few good applications every day.")
    if d["kit_week"] < 5:
        tips.append("Send 5 to 10 personal notes to real recruiters or employees this week.")
    if sum(p for _, p, _ in d["apps"].values()) == 0 and sum(a for a, _, _ in d["apps"].values()) >= 20:
        tips.append("20+ applications with no replies: review the resume against two job descriptions, and push referrals.")
    L.append("Next: " + (" ".join(tips) if tips else "keep the same rhythm."))
    L.append("")
    L.append("Small numbers early on are normal. Judge a source only after about 20 applications.")
    return "\n".join(L)


def main():
    preview, force = "--print" in sys.argv, "--force" in sys.argv
    today = date.today()
    if not preview and not force and os.path.exists(GUARD):
        try:
            last = datetime.fromisoformat(open(GUARD).read().strip()).date()
            if (today - last).days < 7:
                print("[weekly] already sent within the last 7 days; skipping.")
                return 0
        except Exception:
            pass
    body = build_body(collect(today), today)
    if preview or not (EMAIL_ADDRESS and EMAIL_PASSWORD):
        print(body)
        return 0
    msg = MIMEText(body, "plain", "utf-8")
    msg["From"], msg["To"] = EMAIL_ADDRESS, RECIPIENT
    msg["Subject"] = f"[Job Agent] Weekly summary, {today:%d %b}"
    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=30) as s:
            s.ehlo(); s.starttls(); s.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
            s.sendmail(EMAIL_ADDRESS, [RECIPIENT], msg.as_string())
        open(GUARD, "w").write(today.isoformat())
        print("[weekly] summary emailed.")
    except Exception as exc:
        print(f"[weekly] email failed: {str(exc)[:100]}")
    return 0


if __name__ == "__main__":
    main()
