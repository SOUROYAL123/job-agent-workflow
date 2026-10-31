#!/usr/bin/env python3
"""
job_agent.py - resilient multi-portal cloud job discovery, scoring & digest pipeline.

Stages (run all, or one with --stage):
  init     create/verify/migrate SQLite schema (jobs + applications tables)
  scrape   JobSpy boards -> jobs.db
  ats      Greenhouse / Lever / Ashby boards -> jobs.db (needs slugs in .env)
  remote   Worldwide remote feeds (remote_feeds.py) -> jobs.db
  email    Gmail job-alert emails (IMAP) -> jobs.db
  score    Prefilter + HTTP Liveness Gate + Gemini JSON scoring with failover & clone caching
  digest   build Excel backup + email digest of new jobs scoring >= MIN score
  excel    rebuild the Excel backup only (preserving user notes & syncing applications)
  status   print database, applications & match statistics in the terminal

Usage:
  python job_agent.py                           # run full daily pipeline
  python job_agent.py --stage score             # run a single stage
  python job_agent.py --stage status            # inspect DB & top matches
  python job_agent.py --stage score --max 25    # override max API calls for one run
  python job_agent.py --track "Company | Title | applied | Optional note"
"""

import argparse
import email as email_lib
import hashlib
import html
import imaplib
import os
import re
import smtplib
import sqlite3
import sys
import time
import truststore
truststore.inject_into_ssl()
from datetime import datetime, timedelta, timezone
from email import encoders
from email.header import decode_header
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from typing import List, Literal, Optional, Tuple
from urllib.parse import parse_qs, urlparse

from dotenv import load_dotenv
from pydantic import BaseModel, Field

# Ensure Windows PowerShell & Task Scheduler log redirection never crash on Unicode symbols
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

load_dotenv(override=True)  # .env wins over stale system/user environment variables

# ----------------------------------------------------------------------
# CONFIG (all overridable in .env)
# ----------------------------------------------------------------------
DB_PATH = os.getenv("DB_PATH", "jobs.db")
EXCEL_PATH = os.getenv("EXCEL_BACKUP_PATH", "cloud_jobs_backup.xlsx")
PROFILE_PATH = os.getenv("PROFILE_PATH", "candidate_profile.txt")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-flash-lite-latest").strip()
_FALLBACK_RAW = os.getenv(
    "GEMINI_FALLBACK_MODELS",
    "gemini-flash-lite-latest,gemini-3.5-flash-lite,gemini-3.1-flash-lite,gemini-3.5-flash,gemini-3-flash-preview",
)
MODEL_CHAIN: List[str] = []
for _m in [GEMINI_MODEL] + [x.strip() for x in _FALLBACK_RAW.split(",") if x.strip()]:
    if _m and _m not in MODEL_CHAIN:
        MODEL_CHAIN.append(_m)

# Optional free fallback provider (any OpenAI-compatible API: Groq, OpenRouter, local Ollama).
# Used only after every Gemini model is out of quota. Leave FALLBACK_BASE_URL blank to disable.
FALLBACK_BASE_URL = os.getenv("FALLBACK_BASE_URL", "").strip().rstrip("/")
FALLBACK_API_KEY = os.getenv("FALLBACK_API_KEY", "").strip()
FALLBACK_MODELS = [m.strip() for m in os.getenv("FALLBACK_MODELS", "").split(",") if m.strip()]

RERANK_MODEL = os.getenv("RERANK_MODEL", "").strip()   # e.g. gemini-3.5-flash; blank = off
RERANK_FROM = int(os.getenv("RERANK_FROM_SCORE", "50"))
MIN_SCORE = int(os.getenv("MIN_EMAIL_MATCH_SCORE", "60"))
MAX_SCORE_PER_RUN = int(os.getenv("MAX_SCORE_PER_RUN", "60"))
SCORE_SLEEP_SEC = float(os.getenv("SCORE_SLEEP_SEC", "8"))
LIVENESS_CHECK = os.getenv("LIVENESS_CHECK", "1").strip().lower() not in ("0", "false", "no")

IMAP_SERVER = os.getenv("IMAP_SERVER", "imap.gmail.com")
IMAP_PORT = int(os.getenv("IMAP_PORT", "993"))
IMAP_LOOKBACK_DAYS = int(os.getenv("IMAP_LOOKBACK_DAYS", "3"))
SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
EMAIL_ADDRESS = os.getenv("EMAIL_ADDRESS", "").strip()
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "").strip()
RECIPIENT_EMAIL = os.getenv("RECIPIENT_EMAIL", "").strip() or EMAIL_ADDRESS

_proxies = os.getenv("PROXIES", "").strip()
PROXIES = [p.strip() for p in _proxies.split(",") if p.strip()] or None

_DEFAULT_QUERIES = [
    "Cloud Support Engineer AWS",
    "Cloud Support Associate AWS",
    "AWS Implementation Engineer",
    "Technical Support Engineer Cloud SaaS",
    "Cloud Operations NOC Engineer AWS",
    "Junior Cloud Engineer AWS",
]
_env_queries = [q.strip() for q in os.getenv("SEARCH_QUERIES", "").split("|") if q.strip()]
SEARCH_QUERIES = _env_queries or _DEFAULT_QUERIES

_env_locs = [loc.strip() for loc in os.getenv("SEARCH_LOCATIONS", "").split("|") if loc.strip()]
LOCATIONS = _env_locs or ["India", "Kolkata, India"]

_env_sites = [s.strip() for s in os.getenv("SEARCH_SITES", "").split(",") if s.strip()]
SITES = _env_sites or ["linkedin", "indeed", "naukri"]

RESULTS_WANTED = int(os.getenv("RESULTS_WANTED", "12"))
HOURS_OLD = int(os.getenv("HOURS_OLD", "72"))

ATS_GREENHOUSE = [s.strip() for s in os.getenv("ATS_GREENHOUSE", "").split(",") if s.strip()]
ATS_LEVER = [s.strip() for s in os.getenv("ATS_LEVER", "").split(",") if s.strip()]
ATS_ASHBY = [s.strip() for s in os.getenv("ATS_ASHBY", "").split(",") if s.strip()]

# Quota-saving heuristic filters (allow Solutions/Associate Architect & L1/L2 through)
SENIOR_TITLE = re.compile(
    r"\b(senior|sr\.?|lead|principal|staff|manager|head of|director|vp|vice president|avp)\b"
    r"|\b(enterprise|chief|distinguished)\s+architect\b"
    r"|\b(l3|l4|iii|iv)\b",
    re.I,
)
NON_PAID_TITLE = re.compile(
    r"\b(intern|internship|trainee|apprentice|unpaid|volunteer|stipend)\b", re.I)
OFF_DOMAIN_TITLE = re.compile(
    r"\b(data cent(?:er|re)|technician|program associate|sap|abap|peoplesoft|ukg|workday|mainframe|cobol|etrm|salesforce developer|"
    r"ios developer|android developer|react native|angular developer|mlops)\b",
    re.I,
)
HIGH_EXP_FIELD = re.compile(
    r"^\s*(?:[6-9]|1\d)\s*(?:\+|-\s*\d+|\s+to\s+\d+)?\s*(?:years?|yrs?)\b"
    r"|\b(director|executive)\b",
    re.I,
)
HIGH_EXP_DESC = re.compile(
    r"\b(?:minimum|at\s+least|require[sd]?|must\s+have)\s+(?:of\s+)?(?:[6-9]|1\d)\s*\+?\s*"
    r"(?:years?|yrs?)\s+(?:of\s+)?(?:\w+\s+){0,3}experience\b",
    re.I,
)

# Location rule: Kolkata = any work mode (onsite/hybrid/remote); rest of India = remote only
KOLKATA_WORDS = ("kolkata", "calcutta", "west bengal", "howrah", "salt lake", "new town",
                 "rajarhat", "sector v", "barasat", "durgapur", "wb,")
REMOTE_PHRASES = ("remote", "work from home", "work-from-home", "wfh", "anywhere in india")
HYBRID_ONLY = re.compile(r"\bhybrid\b", re.I)

# HTTP Liveness Check patterns (from career-ops)
DEAD_URL_PATTERNS = ("error=true", "job-closed", "expired", "not-found", "job_not_found")
DEAD_PAGE_PATTERNS = re.compile(
    r"(no longer accepting applications|"
    r"this job (posting )?is no longer available|"
    r"position has been filled|"
    r"job you are looking for is no longer open|"
    r"this position is closed|"
    r"sorry, this job is expired|"
    r"job posting has expired)",
    re.I,
)
LIVENESS_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    )
}


def check_job_liveness(url: str, timeout: int = 8) -> Tuple[bool, str]:
    """
    Verifies if a job URL is live before spending Gemini quota.
    Fails open (returns True) on 401/403/429/999 anti-bot walls or network timeouts
    so valid LinkedIn/Indeed/Cloudflare postings are never discarded.
    """
    if not url or not url.startswith("http"):
        return True, "no_http_url"
    try:
        import requests
        resp = requests.get(url, headers=LIVENESS_HEADERS, timeout=timeout, allow_redirects=True)
    except Exception as exc:
        return True, f"network_skip:{type(exc).__name__}"

    if resp.status_code in (404, 410):
        return False, f"HTTP {resp.status_code}"

    if resp.status_code in (401, 403, 429, 999) or resp.status_code >= 500:
        return True, f"bot_or_server_skip:{resp.status_code}"

    final_url = (resp.url or "").lower()
    for pat in DEAD_URL_PATTERNS:
        if pat in final_url:
            return False, f"redirected to {pat}"

    snippet = (resp.text or "")[:25000]
    m = DEAD_PAGE_PATTERNS.search(snippet)
    if m:
        return False, f"closed banner ('{m.group(0).lower()}')"

    return True, "alive"


def location_ok(job: dict) -> bool:
    if job.get("came_from_email"):
        return True  # alert emails carry no real location; let the scorer judge
    loc = (job.get("location") or "").lower().strip()
    if loc in ("", "see listing"):
        return True
    if any(w in loc for w in KOLKATA_WORDS):
        return True
    if job.get("is_remote"):
        return True
    title = (job.get("title") or "").lower()
    if any(w in loc or w in title for w in REMOTE_PHRASES):
        return True
    head = (job.get("description") or "")[:1500].lower()
    if any(p in head for p in ("fully remote", "100% remote", "remote-first", "work from home",
                               "remote position", "remote role", "this is a remote")):
        return True
    return False


ALERT_SENDER_HINTS = (
    "linkedin", "indeed", "naukri", "glassdoor", "hirist", "instahyre",
    "wellfound", "foundit", "cutshort", "timesjobs", "amazon.jobs", "workday",
    "greenhouse", "lever", "ashby",
)
ALERT_SUBJECT_HINTS = (
    "job", "alert", "hiring", "cloud", "aws", "devops", "support engineer",
    "noc", "implementation", "sre",
)
SKIP_ANCHORS = {
    "unsubscribe", "privacy policy", "help", "manage alerts", "settings",
    "view in browser", "see all jobs", "update your preferences", "terms",
    "contact us", "email preferences",
}
GENERIC_ANCHORS = {
    "listing", "view job", "view", "apply", "apply now", "see job", "view details", "job details",
    "see details", "view this job", "view more", "easy apply", "click here", "here", "details",
}


def alert_title(text: str, ctx: str, subject: str) -> str:
    """Anchor text like 'Listing' or 'View job' says nothing; use the surrounding text instead."""
    t = (text or "").strip()
    if t.lower() in GENERIC_ANCHORS or len(t) <= 4:
        c = re.sub(r"\s+", " ", (ctx or "").replace(t, " ")).strip(" -|:\u00b7\u2022")
        if len(c) >= 8:
            return c[:110]
        return f"Email alert role ({subject[:45]})"
    return t[:150]


JOB_LINK_RE = re.compile(
    r"(linkedin\.com/(comm/)?jobs/view|indeed\.com/(rc/clk|viewjob|pagead/clk)|"
    r"naukri\.com/job-listings|glassdoor\.[a-z.]+/(partner/jobListing|job-listing)|"
    r"hirist\.(tech|com)/j/|instahyre\.com/job-|wellfound\.com/jobs|"
    r"foundit\.in/job/|cutshort\.io/job/|timesjobs\.com/job-detail|"
    r"amazon\.jobs/[a-z-]+/jobs/\d+|boards\.greenhouse\.io/[^/]+/jobs/\d+|"
    r"jobs\.lever\.co/[^/]+/[0-9a-f-]{20,}|jobs\.ashbyhq\.com/[^/]+/[0-9a-f-]{16,}|"
    r"myworkdayjobs\.com/.+/job/)",
    re.I,
)


# ----------------------------------------------------------------------
# HELPERS
# ----------------------------------------------------------------------
def now_utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def today_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d")


def is_fresh(date_posted: str, hours_old: int = None) -> bool:
    """Boards give a date, not a time, so allow ceil(hours/24)+0 whole days back. Unknown date = keep."""
    hours_old = HOURS_OLD if hours_old is None else hours_old
    try:
        d = datetime.fromisoformat(str(date_posted).strip()[:10]).date()
    except Exception:
        return True
    max_days = -(-hours_old // 24)  # ceil
    return (datetime.now().date() - d).days <= max_days


def clean(val) -> str:
    if val is None:
        return ""
    try:
        if val != val:  # NaN
            return ""
    except Exception:
        pass
    s = str(val).strip()
    return "" if s.lower() in {"nan", "none", "nat"} else s


def canonical_key(url: str, title: str = "", company: str = "") -> str:
    """Stable identity for a job. Indeed keeps ?jk= id; LinkedIn uses numeric job id."""
    if not url:
        return f"tc:{title.lower()}::{company.lower()}"
    p = urlparse(url)
    host = p.netloc.lower().replace("www.", "")
    qs = parse_qs(p.query)
    if "indeed." in host:
        jk = (qs.get("jk") or qs.get("vjk") or [""])[0]
        if jk:
            return f"indeed:{jk}"
    if "linkedin.com" in host:
        m = re.search(r"/jobs/view/(?:[^/?]*-)?(\d{6,})", p.path)
        if m:
            return f"linkedin:{m.group(1)}"
        cur = (qs.get("currentJobId") or [""])[0]
        if cur:
            return f"linkedin:{cur}"
    if "naukri.com" in host:
        m = re.search(r"-(\d{9,14})(?:\?|$|/)", url)
        if m:
            return f"naukri:{m.group(1)}"
    path = p.path.rstrip("/")
    if not path:
        return f"{host}?{p.query}".lower()
    return f"{host}{path}".lower()


def job_hash(url: str, title: str = "", company: str = "") -> str:
    return hashlib.sha256(canonical_key(url, title, company).encode()).hexdigest()


def clone_signature(job: dict) -> Optional[str]:
    """Fingerprint identical multi-city postings so Gemini only scores each unique JD once."""
    desc = re.sub(r"\s+", " ", (job.get("description") or "")).strip().lower()
    if len(desc) < 120:
        return None
    title = re.sub(r"\s+", " ", (job.get("title") or "")).strip().lower()
    company = re.sub(r"\s+", " ", (job.get("company") or "")).strip().lower()
    raw = f"{title}||{company}||{desc[:500]}"
    return hashlib.md5(raw.encode("utf-8", errors="ignore")).hexdigest()


def get_conn(db_path: str = DB_PATH) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


JOB_COLS = (
    "job_hash", "title", "company", "location", "experience", "job_type", "is_remote",
    "date_posted", "description", "job_url", "job_url_direct", "company_url",
    "source_channel", "source_site", "came_from_email",
    "source_email_sender", "source_email_subject", "source_email_date",
)


# ----------------------------------------------------------------------
# APPLIED-JOB MEMORY & `applications` TABLE
# ----------------------------------------------------------------------
APPLIED_NOTE_RE = re.compile(os.getenv("APPLIED_NOTE_REGEX", r"appl|interview|reject|offer|skip"), re.I)

_COMPANY_NOISE = {
    "pvt", "ltd", "private", "limited", "inc", "llc", "llp", "corp", "corporation", "co",
    "india", "the", "technologies", "technology", "solutions", "services", "systems", "global",
}
_TITLE_NOISE = {
    "remote", "hybrid", "onsite", "wfh", "india", "kolkata", "calcutta", "bengaluru", "bangalore",
    "hyderabad", "pune", "mumbai", "chennai", "delhi", "ncr", "gurgaon", "gurugram", "noida",
    "urgent", "hiring", "opening", "immediate", "joiner", "joiners",
}


def _norm_tokens(text: str, noise: set) -> str:
    text = re.sub(r"[\(\[\{][^\)\]\}]*[\)\]\}]", " ", (text or "").lower())  # drop (Remote), [Kolkata]
    text = re.sub(r"[^a-z0-9+#]+", " ", text)
    return " ".join(t for t in text.split() if t not in noise)


def applied_key(title: str, company: str) -> Optional[str]:
    """Site-independent identity: normalised company + title. None when the company is unknown."""
    comp_raw = (company or "").strip().lower()
    if not comp_raw or comp_raw.startswith("via ") or comp_raw == "unknown company":
        return None
    comp = _norm_tokens(company, _COMPANY_NOISE)
    ttl = _norm_tokens(title, _TITLE_NOISE)
    if not comp or not ttl:
        return None
    return f"{comp}||{ttl}"


def note_is_applied(note: str) -> bool:
    return bool(note and note.strip() and note.strip() != "-" and APPLIED_NOTE_RE.search(note))


def _infer_status_from_note(note: str) -> str:
    n = (note or "").lower()
    if "offer" in n:
        return "offer"
    if "interview" in n:
        return "interviewing"
    if any(w in n for w in ("reject", "declin", "not selected")):
        return "rejected"
    if "skip" in n:
        return "skip"
    return "applied"


def sync_applications_from_jobs(conn: sqlite3.Connection) -> int:
    """
    Ensures any row in `jobs` with an application note in `user_notes`
    is mirrored into the structured `applications` table.
    """
    cur = conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='applications'")
    if not cur.fetchone():
        return 0

    rows = conn.execute(
        "SELECT title, company, COALESCE(NULLIF(job_url_direct, ''), job_url) AS url, user_notes "
        "FROM jobs WHERE TRIM(COALESCE(user_notes, '')) != ''"
    ).fetchall()
    synced = 0
    today = today_iso()
    default_fup = (datetime.now(timezone.utc) + timedelta(days=6)).strftime("%Y-%m-%d")

    for r in rows:
        note = (r["user_notes"] or "").strip()
        if not note_is_applied(note):
            continue
        k = applied_key(r["title"], r["company"])
        if not k:
            continue
        status = _infer_status_from_note(note)
        fup = default_fup if status == "applied" else ""
        conn.execute(
            """
            INSERT INTO applications (app_key, company, title, job_url, status, applied_date, follow_up_date, notes, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(app_key) DO UPDATE SET
                status = excluded.status,
                notes = excluded.notes,
                job_url = COALESCE(NULLIF(excluded.job_url, ''), applications.job_url),
                updated_at = excluded.updated_at
            """,
            (k, r["company"] or "", r["title"] or "", r["url"] or "", status, today, fup, note, now_utc()),
        )
        synced += 1
    conn.commit()
    return synced


def upsert_application(
    db_path: str,
    company: str,
    title: str,
    status: str = "applied",
    job_url: str = "",
    follow_up_date: str = "",
    notes: str = "",
) -> Optional[str]:
    """Directly insert or update a tracked application in the `applications` table."""
    stage_init(db_path)
    k = applied_key(title, company)
    if not k:
        k = f"{(company or 'unknown').strip().lower()}||{(title or 'untitled').strip().lower()}"
    today = today_iso()
    if not follow_up_date and status.lower() == "applied":
        follow_up_date = (datetime.now(timezone.utc) + timedelta(days=6)).strftime("%Y-%m-%d")

    conn = get_conn(db_path)
    conn.execute(
        """
        INSERT INTO applications (app_key, company, title, job_url, status, applied_date, follow_up_date, notes, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(app_key) DO UPDATE SET
            status = excluded.status,
            job_url = COALESCE(NULLIF(excluded.job_url, ''), applications.job_url),
            follow_up_date = COALESCE(NULLIF(excluded.follow_up_date, ''), applications.follow_up_date),
            notes = COALESCE(NULLIF(excluded.notes, ''), applications.notes),
            updated_at = excluded.updated_at
        """,
        (k, company.strip(), title.strip(), job_url.strip(), status.strip().lower(), today, follow_up_date, notes.strip(), now_utc()),
    )
    conn.commit()
    conn.close()
    return k


def load_applied_keys(conn: sqlite3.Connection) -> set:
    keys = set()
    for r in conn.execute("SELECT title, company, user_notes FROM jobs "
                          "WHERE TRIM(COALESCE(user_notes,'')) != ''"):
        if note_is_applied(r["user_notes"]):
            k = applied_key(r["title"], r["company"])
            if k:
                keys.add(k)
    try:
        for r in conn.execute(
            "SELECT app_key, title, company FROM applications "
            "WHERE status IN ('applied', 'interviewing', 'rejected', 'offer', 'skip')"
        ):
            if r["app_key"]:
                keys.add(r["app_key"])
            k2 = applied_key(r["title"], r["company"])
            if k2:
                keys.add(k2)
    except sqlite3.Error:
        pass
    return keys


def is_applied_duplicate(job: dict, applied_keys: set) -> bool:
    """True if this row is itself marked applied, or matches an applied job from any site."""
    if note_is_applied(job.get("user_notes") or ""):
        return True
    k = applied_key(job.get("title"), job.get("company"))
    return bool(k and k in applied_keys)


_APPLIED_CACHE: dict = {"keys": set()}


def refresh_applied_cache(db_path: str = DB_PATH) -> int:
    """Reload applied keys so insert_job can skip them. Called at the start of each inserting stage."""
    try:
        conn = get_conn(db_path)
        try:
            _APPLIED_CACHE["keys"] = load_applied_keys(conn)
        finally:
            conn.close()
    except sqlite3.Error:
        _APPLIED_CACHE["keys"] = set()
    return len(_APPLIED_CACHE["keys"])


def insert_job(cur: sqlite3.Cursor, **f) -> bool:
    k = applied_key(f.get("title", ""), f.get("company", ""))
    if k and k in _APPLIED_CACHE["keys"]:
        return False  # same company+title already applied on another site
    vals = [f.get(c, "" if c not in ("is_remote", "came_from_email") else 0) for c in JOB_COLS]
    cur.execute(
        f"INSERT OR IGNORE INTO jobs ({','.join(JOB_COLS)}) VALUES ({','.join('?' * len(JOB_COLS))})",
        vals,
    )
    return cur.rowcount > 0


# ----------------------------------------------------------------------
# STAGE: INIT
# ----------------------------------------------------------------------
def stage_init(db_path: str = DB_PATH) -> None:
    conn = get_conn(db_path)
    conn.execute("""
    CREATE TABLE IF NOT EXISTS jobs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        job_hash TEXT UNIQUE NOT NULL,
        title TEXT NOT NULL, company TEXT, location TEXT, experience TEXT, job_type TEXT,
        is_remote INTEGER DEFAULT 0, date_posted TEXT, description TEXT,
        job_url TEXT NOT NULL, job_url_direct TEXT, company_url TEXT,
        source_channel TEXT NOT NULL, source_site TEXT,
        came_from_email INTEGER DEFAULT 0,
        source_email_sender TEXT, source_email_subject TEXT, source_email_date TEXT,
        scored INTEGER DEFAULT 0, match_score INTEGER, priority TEXT, fit_reason TEXT,
        missing_skills TEXT, recommendation TEXT, resume_tip TEXT DEFAULT '',
        scored_model TEXT DEFAULT '', scored_at TEXT,
        emailed_in_digest INTEGER DEFAULT 0, emailed_at TEXT,
        user_notes TEXT DEFAULT '',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")
    cols = {r[1] for r in conn.execute("PRAGMA table_info(jobs)")}
    for col_name in ("user_notes", "resume_tip", "scored_model"):
        if col_name not in cols:
            conn.execute(f"ALTER TABLE jobs ADD COLUMN {col_name} TEXT DEFAULT ''")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_scored ON jobs(scored)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_jobs_emailed ON jobs(emailed_in_digest, match_score)")

    # New structured `applications` table (from career-ops architecture)
    conn.execute("""
    CREATE TABLE IF NOT EXISTS applications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        app_key TEXT UNIQUE NOT NULL,
        company TEXT NOT NULL,
        title TEXT NOT NULL,
        job_url TEXT DEFAULT '',
        status TEXT NOT NULL DEFAULT 'applied',
        applied_date TEXT NOT NULL,
        follow_up_date TEXT DEFAULT '',
        notes TEXT DEFAULT '',
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_applications_key ON applications(app_key)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_applications_status ON applications(status)")
    sync_applications_from_jobs(conn)

    conn.commit()
    conn.close()
    print(f"[init] database ready at {db_path} (jobs + applications tables verified)")


# ----------------------------------------------------------------------
# STAGE: SCRAPE (JobSpy)
# ----------------------------------------------------------------------
def stage_scrape(db_path: str = DB_PATH) -> int:
    try:
        from jobspy import scrape_jobs
    except ImportError:
        print("[scrape] python-jobspy not installed; skipping.")
        return 0

    import inspect

    # JobSpy renamed linkedin_fetch_description -> fetch_description; use whichever this install has.
    try:
        _params = inspect.signature(scrape_jobs).parameters
    except (TypeError, ValueError):
        _params = {}
    if "fetch_description" in _params:
        desc_kw = {"fetch_description": True}
    elif "linkedin_fetch_description" in _params or not _params:
        desc_kw = {"linkedin_fetch_description": True}
    else:
        desc_kw = {}

    n_applied = refresh_applied_cache(db_path)
    if n_applied:
        print(f"[scrape] {n_applied} applied job(s) on file; same company+title will be skipped on any site")

    conn = get_conn(db_path)
    cur = conn.cursor()
    added = 0
    for query in SEARCH_QUERIES:
        for location in LOCATIONS:
            print(f"[scrape] '{query}' / '{location}' on {SITES}")
            india_wide = location.strip().lower() == "india"  # outside Kolkata we only want remote roles
            other_sites = [s for s in SITES if s.lower() != "indeed"]
            has_indeed = any(s.lower() == "indeed" for s in SITES)
            base = dict(search_term=query, location=location, results_wanted=RESULTS_WANTED,
                        country_indeed="India", verbose=0, **desc_kw)
            if PROXIES:
                base["proxies"] = PROXIES

            # Indeed accepts only ONE of: hours_old | (job_type + is_remote) | easy_apply per search,
            # so it gets its own calls. LinkedIn/Naukri etc. keep both filters together.
            calls = []
            if other_sites:
                kw = dict(base, site_name=other_sites, hours_old=HOURS_OLD)
                if india_wide:
                    kw["is_remote"] = True
                calls.append(("other", kw))
            if has_indeed:
                calls.append(("indeed", dict(base, site_name=["indeed"], hours_old=HOURS_OLD)))
                if india_wide:
                    # remote-only pass, no hours_old; freshness is enforced locally by is_fresh()
                    calls.append(("indeed-remote", dict(base, site_name=["indeed"], is_remote=True)))

            for label, kwargs in calls:
                try:
                    df = scrape_jobs(**kwargs)
                except Exception as exc:
                    print(f"[scrape] WARNING {query}/{location} ({label}): {exc}")
                    continue
                if df is None or df.empty:
                    continue
                batch_added = stale = 0
                for _, row in df.iterrows():
                    url = clean(row.get("job_url"))
                    if not url:
                        continue
                    if not is_fresh(clean(row.get("date_posted"))):
                        stale += 1
                        continue
                    title = clean(row.get("title")) or "Untitled role"
                    company = clean(row.get("company")) or "Unknown company"
                    ok = insert_job(
                        cur,
                        job_hash=job_hash(url, title, company), title=title, company=company,
                        location=clean(row.get("location")) or location,
                        experience=clean(row.get("job_level")) or clean(row.get("experience_range")),
                        job_type=clean(row.get("job_type")),
                        is_remote=1 if row.get("is_remote") is True else 0,
                        date_posted=clean(row.get("date_posted")),
                        description=clean(row.get("description")),
                        job_url=url, job_url_direct=clean(row.get("job_url_direct")) or url,
                        company_url=clean(row.get("company_url")),
                        source_channel="JobSpy Scraper", source_site=clean(row.get("site")) or "jobspy",
                    )
                    batch_added += int(ok)
                added += batch_added
                if stale:
                    print(f"[scrape]   {label}: skipped {stale} older than {HOURS_OLD}h")
                conn.commit()
                time.sleep(3)
    conn.close()
    print(f"[scrape] {added} new job(s)")
    return added


# ----------------------------------------------------------------------
# STAGE: ATS (Greenhouse / Lever / Ashby direct feeds)
# ----------------------------------------------------------------------
INDIA_WORDS = (
    "india", "bengaluru", "bangalore", "hyderabad", "pune", "mumbai", "gurgaon",
    "gurugram", "noida", "chennai", "kolkata", "remote", "apac", "anywhere",
)
ROLE_WORDS = (
    "aws", "cloud", "support", "implementation", "devops", "operations", "sre",
    "noc", "infrastructure", "onboarding", "platform", "reliability",
)


def stage_ats(db_path: str = DB_PATH) -> int:
    if not (ATS_GREENHOUSE or ATS_LEVER or ATS_ASHBY):
        print("[ats] no ATS_GREENHOUSE / ATS_LEVER / ATS_ASHBY slugs set; skipping.")
        return 0
    import requests
    from bs4 import BeautifulSoup

    def relevant(title: str, loc: str) -> bool:
        t, l = title.lower(), loc.lower()
        return any(w in t for w in ROLE_WORDS) and any(w in l for w in INDIA_WORDS)

    refresh_applied_cache(db_path)
    conn = get_conn(db_path)
    cur = conn.cursor()
    added = 0

    # 1. Greenhouse
    for slug in ATS_GREENHOUSE:
        try:
            r = requests.get(
                f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs?content=true",
                timeout=15,
            )
            if r.status_code != 200:
                print(f"[ats] greenhouse '{slug}' -> HTTP {r.status_code} (check slug)")
                continue
            for j in r.json().get("jobs", []):
                title, loc = j.get("title", ""), (j.get("location") or {}).get("name", "")
                url = j.get("absolute_url", "")
                if not url or not relevant(title, loc):
                    continue
                desc = BeautifulSoup(html.unescape(j.get("content", "")), "html.parser").get_text(" ", strip=True)
                added += int(insert_job(
                    cur, job_hash=job_hash(url), title=title, company=slug.title(), location=loc,
                    is_remote=1 if "remote" in loc.lower() else 0,
                    description=desc, job_url=url, job_url_direct=url,
                    company_url=f"https://boards.greenhouse.io/{slug}",
                    date_posted=j.get("updated_at", ""),
                    source_channel="Direct ATS", source_site="greenhouse",
                ))
        except Exception as exc:
            print(f"[ats] greenhouse '{slug}': {exc}")

    # 2. Lever
    for slug in ATS_LEVER:
        try:
            r = requests.get(f"https://api.lever.co/v0/postings/{slug}?mode=json", timeout=15)
            if r.status_code != 200:
                print(f"[ats] lever '{slug}' -> HTTP {r.status_code} (check slug)")
                continue
            for j in r.json():
                title = j.get("text", "")
                loc = (j.get("categories") or {}).get("location", "") or ""
                url = j.get("hostedUrl", "")
                if not url or not relevant(title, loc):
                    continue
                added += int(insert_job(
                    cur, job_hash=job_hash(url), title=title, company=slug.title(), location=loc,
                    is_remote=1 if "remote" in loc.lower() else 0,
                    description=j.get("descriptionPlain", ""), job_url=url,
                    job_url_direct=j.get("applyUrl", url),
                    company_url=f"https://jobs.lever.co/{slug}",
                    source_channel="Direct ATS", source_site="lever",
                ))
        except Exception as exc:
            print(f"[ats] lever '{slug}': {exc}")

    # 3. Ashby HQ (public JSON posting API)
    for slug in ATS_ASHBY:
        try:
            r = requests.get(
                f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=true",
                timeout=15,
            )
            if r.status_code != 200:
                print(f"[ats] ashby '{slug}' -> HTTP {r.status_code} (check slug)")
                continue
            for j in r.json().get("jobs", []):
                if not j.get("isListed", True):
                    continue
                title = (j.get("title") or "").strip()
                sec_locs = [
                    loc_obj.get("location", "")
                    for loc_obj in (j.get("secondaryLocations") or [])
                    if isinstance(loc_obj, dict) and loc_obj.get("location")
                ]
                primary_loc = j.get("location") or ""
                loc = ", ".join(filter(None, [primary_loc, *sec_locs]))
                is_rem = bool(j.get("isRemote")) or "remote" in loc.lower()
                if is_rem and "remote" not in loc.lower():
                    loc = f"{loc} (Remote)".strip()
                url = j.get("jobUrl") or j.get("applyUrl") or ""
                if not url or not relevant(title, loc):
                    continue
                desc = j.get("descriptionPlain") or ""
                if not desc and j.get("descriptionHtml"):
                    desc = BeautifulSoup(html.unescape(j["descriptionHtml"]), "html.parser").get_text(" ", strip=True)
                added += int(insert_job(
                    cur, job_hash=job_hash(url), title=title,
                    company=slug.replace("-", " ").title(), location=loc,
                    is_remote=1 if is_rem else 0,
                    date_posted=j.get("publishedAt", ""),
                    description=desc, job_url=url,
                    job_url_direct=j.get("applyUrl") or url,
                    company_url=f"https://jobs.ashbyhq.com/{slug}",
                    source_channel="Direct ATS", source_site="ashby",
                ))
        except Exception as exc:
            print(f"[ats] ashby '{slug}': {exc}")

    conn.commit()
    conn.close()
    print(f"[ats] {added} new job(s)")
    return added


# ----------------------------------------------------------------------
# STAGE: EMAIL (IMAP)
# ----------------------------------------------------------------------
def decode_mime(val: str) -> str:
    if not val:
        return ""
    out = []
    for part, enc in decode_header(val):
        out.append(part.decode(enc or "utf-8", errors="replace") if isinstance(part, bytes) else str(part))
    return "".join(out).strip()


def email_body(msg) -> str:
    html_part, text_part = "", ""
    for part in (msg.walk() if msg.is_multipart() else [msg]):
        if "attachment" in str(part.get("Content-Disposition") or ""):
            continue
        ctype = part.get_content_type()
        if ctype not in ("text/html", "text/plain"):
            continue
        payload = part.get_payload(decode=True)
        if not payload:
            continue
        text = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
        if ctype == "text/html":
            html_part += text
        else:
            text_part += text
    return html_part or text_part


def parse_alert_message(msg) -> List[dict]:
    """Pure function: email.Message -> list of job dicts. Testable offline without Gmail."""
    from bs4 import BeautifulSoup

    subject = decode_mime(msg.get("Subject", ""))
    sender = decode_mime(msg.get("From", ""))
    date = decode_mime(msg.get("Date", ""))
    s_low, sub_low = sender.lower(), subject.lower()
    if not (any(h in s_low for h in ALERT_SENDER_HINTS) or any(h in sub_low for h in ALERT_SUBJECT_HINTS)):
        return []
    body = email_body(msg)
    if not body:
        return []
    soup = BeautifulSoup(body, "html.parser")
    summary = soup.get_text(" ", strip=True)[:3000]
    jobs, seen = [], set()
    for a in soup.find_all("a", href=True):
        href, text = a["href"].strip(), a.get_text(" ", strip=True)
        if not href.startswith("http") or text.lower() in SKIP_ANCHORS:
            continue
        if not JOB_LINK_RE.search(href):
            continue
        key = job_hash(href)
        if key in seen:
            continue
        seen.add(key)
        domain = urlparse(href).netloc.replace("www.", "")
        ctx = a.parent.get_text(" ", strip=True) if a.parent else ""
        jobs.append(dict(
            job_hash=key,
            title=alert_title(text, ctx, subject),
            company=f"Via {domain}", location="See listing", experience="See listing",
            job_type="", is_remote=1 if "remote" in ctx.lower() else 0, date_posted=date,
            description=(f"From email alert.\nSubject: {subject}\nSender: {sender}\n"
                         f"Context: {ctx}\n\nEmail text:\n{summary}"),
            job_url=href, job_url_direct=href, company_url=f"https://{domain}",
            source_channel="Email Alert (IMAP)", source_site=domain, came_from_email=1,
            source_email_sender=sender, source_email_subject=subject, source_email_date=date,
        ))
    return jobs


def stage_email(db_path: str = DB_PATH) -> int:
    if not (EMAIL_ADDRESS and EMAIL_PASSWORD):
        print("[email] EMAIL_ADDRESS/EMAIL_PASSWORD not set; skipping.")
        return 0
    refresh_applied_cache(db_path)
    conn = get_conn(db_path)
    cur = conn.cursor()
    added = 0
    try:
        mail = imaplib.IMAP4_SSL(IMAP_SERVER, IMAP_PORT)
        mail.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
        mail.select("INBOX")
        since = (datetime.now() - timedelta(days=IMAP_LOOKBACK_DAYS)).strftime("%d-%b-%Y")
        status, data = mail.search(None, f"(UNSEEN SINCE {since})")
        ids = data[0].split() if status == "OK" and data and data[0] else []
        print(f"[email] {len(ids)} unread message(s) in last {IMAP_LOOKBACK_DAYS} day(s)")
        for e_id in ids:
            res, msg_data = mail.fetch(e_id, "(BODY.PEEK[])")
            if res != "OK":
                continue
            for part in msg_data:
                if not isinstance(part, tuple):
                    continue
                found = parse_alert_message(email_lib.message_from_bytes(part[1]))
                for job in found:
                    added += int(insert_job(cur, **job))
                if found:
                    mail.store(e_id, "+FLAGS", "\\Seen")
        conn.commit()
        mail.logout()
    except Exception as exc:
        print(f"[email] WARNING: {exc}")
    finally:
        conn.close()
    print(f"[email] {added} new job(s) from alerts")
    return added


# ----------------------------------------------------------------------
# STAGE: SCORE (Prefilter + Liveness Gate + Gemini Evidence Scorer)
# ----------------------------------------------------------------------
class JobMatch(BaseModel):
    match_score: int = Field(description="Fit score 0-100 for this candidate.")
    priority: Literal["High", "Medium", "Low"]
    fit_reason: str = Field(
        description="1-2 sentences explaining fit, quoting a short key phrase from the JD as evidence."
    )
    missing_skills: List[str] = Field(description="Required skills from the JD that the candidate lacks.")
    recommendation: Literal["Apply now", "Maybe", "Skip"]
    resume_tip: str = Field(
        default="",
        description="If match_score >= 55, one concise tip on which candidate project/skill to highlight; else empty string.",
    )


DEFAULT_PROFILE = """\
EDIT candidate_profile.txt - this is only a placeholder.
Entry-level cloud candidate, no professional cloud job yet. Studying AWS SAA-C03."""


def load_profile() -> str:
    if os.path.exists(PROFILE_PATH):
        with open(PROFILE_PATH, encoding="utf-8") as f:
            return f.read().strip()
    return DEFAULT_PROFILE


def build_prompt(job: dict, profile: str) -> str:
    return f"""You are an accurate, honest technical recruiter screening jobs for an entry-level / career-transitioning cloud candidate in India.
Evaluate the job strictly against the CANDIDATE PROFILE below using a two-stage gate + evidence rubric:

STAGE 1 — HARD DEALBREAKER GATE (Score 0-25 / Low / Skip if violated):
- Explicitly requires US/EU/non-Indian citizenship or security clearance with no India option.
- Explicitly requires 5+ years of mandatory corporate DevOps/Cloud production experience.
- On-site/Hybrid role outside Kolkata/West Bengal (unless the JD allows remote/WFH in India).
Note: Silence in a short JD is NEVER a dealbreaker.

STAGE 2 — ANCHORED SCORING RUBRIC:
- 80-100 (High / Apply now): Strong match for 0-2 yr Cloud Support (L1/L2), AWS Support/Associate, CloudOps/NOC/GOC, Implementation Specialist, or Cloud/SaaS Technical Support using AWS, Linux, SQL/PostgreSQL, networking, or API/webhook troubleshooting.
- 60-79 (Medium / Apply now or Maybe): Worth applying. Core responsibilities align (support, operations, implementation, junior SRE/cloud), even if 1-2 secondary tools are missing or experience asks ~2-3 years.
- 40-59 (Low / Maybe): Partial overlap (e.g., Windows-heavy L2 support, Azure/GCP-only focus, or 3-4 yrs experience with steep tooling gaps).
- 0-39 (Low / Skip): Fails Stage 1 Gate, requires expert production Kubernetes/OpenShift/MLOps, pure software engineering from scratch, or unrelated domains.

STRICT EVIDENCE RULES:
1. Credit ONLY skills, labs, degrees, and live self-employed experience explicitly written in the profile. Never invent skills.
2. In `fit_reason`, include a short direct quote from the JD (in single quotes) showing the core requirement that drove your score.
3. Fairly value the candidate's live production webhook/PostgreSQL system (LegacyLens), ap-south-1 2-AZ Terraform/ECS/RDS lab, Linux/network diagnostics, B.Sc. degree, and 14 years of client-facing delivery when evaluating support, NOC/GOC, and implementation roles.
4. If the description is only a short email alert snippet with sparse detail, score conservatively on title/company/location and state that in `fit_reason` without guessing hidden requirements.
5. Populate `resume_tip` with a 1-sentence pitch angle when match_score >= 55 (otherwise "").

CANDIDATE PROFILE:
{profile}

JOB LISTING:
Title: {job.get('title', '')}
Company: {job.get('company', '')}
Location: {job.get('location', '')}
Experience: {job.get('experience', '')}
Remote: {'Yes' if job.get('is_remote') else 'Unspecified'}
Source: {job.get('source_channel', '')} ({job.get('source_site', '')})
Description:
{(job.get('description') or '')[:6500]}"""


def score_single_job(client, job: dict, profile: str, model_name: str = None) -> JobMatch:
    from google.genai import types

    resp = client.models.generate_content(
        model=model_name or GEMINI_MODEL,
        contents=build_prompt(job, profile),
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=JobMatch,
            temperature=0.15,
        ),
    )
    result = JobMatch.model_validate_json(resp.text)
    result.match_score = max(0, min(100, result.match_score))
    return result


def score_via_openai_compat(job: dict, profile: str, base_url: str, api_key: str, model: str) -> JobMatch:
    """Score one job through any OpenAI-compatible /chat/completions endpoint."""
    import requests

    schema_hint = (
        "\n\nReply with ONLY one JSON object, no markdown, with exactly these keys: "
        '"match_score" (integer 0-100), "priority" ("High"|"Medium"|"Low"), "fit_reason" (string), '
        '"missing_skills" (array of strings), "recommendation" ("Apply now"|"Maybe"|"Skip"), '
        '"resume_tip" (string, empty if match_score < 55).'
    )
    headers = {"Content-Type": "application/json"}
    if api_key:
        headers["Authorization"] = f"Bearer {api_key}"
    body = {
        "model": model,
        "messages": [{"role": "user", "content": build_prompt(job, profile) + schema_hint}],
        "temperature": 0.15,
        "response_format": {"type": "json_object"},
    }
    r = requests.post(f"{base_url}/chat/completions", headers=headers, json=body, timeout=90)
    if r.status_code != 200:
        raise RuntimeError(f"{r.status_code} fallback provider: {r.text[:200]}")
    text = r.json()["choices"][0]["message"]["content"] or ""
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
    m = re.search(r"\{.*\}", text, re.S)
    result = JobMatch.model_validate_json(m.group(0) if m else text)
    result.match_score = max(0, min(100, result.match_score))
    return result


def _is_auth_error(exc: Exception) -> bool:
    s = str(exc)
    return (
        "API_KEY_INVALID" in s or "API key not valid" in s or "PERMISSION_DENIED" in s
        or "UNAUTHENTICATED" in s or s.startswith("401") or s.startswith("403")
    )


def _is_quota_error(exc: Exception) -> bool:
    s = str(exc)
    return any(
        tok in s for tok in ("429", "RESOURCE_EXHAUSTED", "503", "UNAVAILABLE", "DEADLINE_EXCEEDED")
    ) or "quota" in s.lower() or "overloaded" in s.lower()


def _parse_retry_seconds(exc: Exception) -> Optional[int]:
    s = str(exc)
    m_hr = re.search(r"retry in\s*(\d+)h", s, re.I)
    if m_hr:
        return int(m_hr.group(1)) * 3600
    m_delay = re.search(r"retryDelay['\"]?\s*:\s*['\"](\d+)s['\"]", s)
    if m_delay:
        return int(m_delay.group(1))
    m_sec = re.search(r"retry in\s*(\d+(?:\.\d+)?)s", s, re.I)
    if m_sec:
        return int(float(m_sec.group(1))) + 1
    return None


def _is_daily_quota_error(exc: Exception) -> bool:
    s = str(exc)
    if "PerDay" in s or "requests per day" in s.lower():
        return True
    delay = _parse_retry_seconds(exc)
    return delay is not None and delay >= 300


def _prefilter_reason(job: dict) -> Optional[str]:
    title = job.get("title") or ""
    exp = job.get("experience") or ""
    desc = job.get("description") or ""
    if SENIOR_TITLE.search(title):
        return "Auto-skipped: senior/lead/L3+ title."
    if NON_PAID_TITLE.search(title) or NON_PAID_TITLE.search(exp):
        return "Auto-skipped: intern/trainee/unpaid role (profile wants paid full-time)."
    if OFF_DOMAIN_TITLE.search(title):
        return "Auto-skipped: off-domain role title."
    if not location_ok(job):
        return "Auto-skipped: outside Kolkata and not remote."
    if HIGH_EXP_FIELD.search(exp) or HIGH_EXP_DESC.search(desc[:3000]):
        return "Auto-skipped: requires 6+ years experience."
    return None


def _load_clone_cache(conn: sqlite3.Connection) -> dict:
    cache = {}
    rows = conn.execute(
        "SELECT title, company, description, match_score, priority, fit_reason, "
        "missing_skills, recommendation, resume_tip FROM jobs WHERE scored = 1 AND match_score > 0"
    )
    for r in rows:
        d = dict(r)
        sig = clone_signature(d)
        if sig and sig not in cache:
            skills = [s.strip() for s in (d.get("missing_skills") or "").split(",") if s.strip()]
            try:
                cache[sig] = JobMatch(
                    match_score=int(d["match_score"] or 0),
                    priority=d["priority"] or "Low",
                    fit_reason=d["fit_reason"] or "",
                    missing_skills=skills,
                    recommendation=d["recommendation"] or "Skip",
                    resume_tip=d.get("resume_tip") or "",
                )
            except Exception:
                continue
    return cache


def stage_score(
    db_path: str = DB_PATH,
    scorer=None,
    sleep_sec: float = None,
    max_calls: int = None,
) -> int:
    """Score unscored rows. `scorer(job_dict) -> JobMatch` can be injected for offline tests."""
    sleep_sec = SCORE_SLEEP_SEC if sleep_sec is None else sleep_sec
    limit_calls = MAX_SCORE_PER_RUN if max_calls is None else max_calls

    conn = get_conn(db_path)
    cur = conn.cursor()
    rows = [dict(r) for r in cur.execute("SELECT * FROM jobs WHERE scored = 0 ORDER BY id DESC")]
    print(f"[score] {len(rows)} unscored job(s)")
    if not rows:
        conn.close()
        return 0

    custom_scorer = scorer is not None
    active_models = list(MODEL_CHAIN) if not custom_scorer else ["custom"]
    clone_cache = {} if custom_scorer else _load_clone_cache(conn)

    client = None
    profile = ""
    if not custom_scorer:
        if not GEMINI_API_KEY:
            conn.close()
            raise SystemExit("[score] GEMINI_API_KEY missing in .env")
        from google import genai
        client = genai.Client(api_key=GEMINI_API_KEY)
        profile = load_profile()
        print(f"[score] model chain: {' -> '.join(active_models)}")

    done = api_calls = 0
    rerank_on = bool(RERANK_MODEL)
    fb_models = list(FALLBACK_MODELS) if (FALLBACK_BASE_URL and not custom_scorer) else []
    fb_active = False
    applied_keys = load_applied_keys(conn)

    for job in rows:
        skip_why = (
            "Auto-skipped: already applied (same company + title marked applied, any site)."
            if is_applied_duplicate(job, applied_keys) else None
        ) or _prefilter_reason(job)
        if skip_why:
            cur.execute(
                "UPDATE jobs SET scored=1, match_score=0, priority='Low', recommendation='Skip', "
                "fit_reason=?, missing_skills='', resume_tip='', scored_model='prefilter', "
                "scored_at=? WHERE id=?",
                (skip_why, now_utc(), job["id"]),
            )
            conn.commit()
            done += 1
            continue

        sig = None if custom_scorer else clone_signature(job)
        if sig and sig in clone_cache:
            cached = clone_cache[sig]
            cur.execute(
                "UPDATE jobs SET scored=1, match_score=?, priority=?, fit_reason=?, "
                "missing_skills=?, recommendation=?, resume_tip=?, scored_model='clone-cache', "
                "scored_at=? WHERE id=?",
                (cached.match_score, cached.priority, cached.fit_reason,
                 ", ".join(cached.missing_skills), cached.recommendation, cached.resume_tip,
                 now_utc(), job["id"]),
            )
            conn.commit()
            done += 1
            print(f"  {cached.match_score:>3}  {job['title'][:46]} @ {job['company'][:22]} [clone]")
            continue

        # HTTP Liveness Check (skips dead 404/410/closed postings before spending LLM quota)
        if not custom_scorer and LIVENESS_CHECK:
            target_url = job.get("job_url_direct") or job.get("job_url") or ""
            alive, live_reason = check_job_liveness(target_url)
            if not alive:
                dead_msg = f"Auto-skipped: dead listing ({live_reason})."
                cur.execute(
                    "UPDATE jobs SET scored=1, match_score=0, priority='Low', recommendation='Skip', "
                    "fit_reason=?, missing_skills='', resume_tip='', scored_model='liveness-gate', "
                    "scored_at=? WHERE id=?",
                    (dead_msg, now_utc(), job["id"]),
                )
                conn.commit()
                done += 1
                print(f"    0  {job['title'][:46]} @ {job['company'][:22]} [dead: {live_reason}]")
                continue

        if api_calls >= limit_calls:
            print(f"[score] hit MAX_SCORE_PER_RUN={limit_calls}; remaining rows wait for next run")
            break

        result = None
        used_model = active_models[0] if active_models else ""
        exhausted_all = False

        while fb_active and fb_models and result is None:
            try:
                api_calls += 1
                result = score_via_openai_compat(job, profile, FALLBACK_BASE_URL, FALLBACK_API_KEY, fb_models[0])
                used_model = f"fallback:{fb_models[0]}"
            except Exception as exc:
                print(f"[score] fallback '{fb_models[0]}' failed ({str(exc)[:100]}); trying next fallback model")
                fb_models.pop(0)
        if fb_active and result is None and not fb_models:
            print("[score] QUOTA_EXHAUSTED across all configured models; stopping cleanly. Unscored rows stay for next run.")
            conn.close()
            return done

        while active_models and result is None and not fb_active:
            current_model = active_models[0]
            rotate_model = False

            for attempt in range(3):
                try:
                    api_calls += 1
                    if custom_scorer:
                        result = scorer(job)
                    else:
                        result = score_single_job(client, job, profile, model_name=current_model)
                    used_model = current_model
                    break
                except Exception as exc:
                    if _is_auth_error(exc):
                        print("[score] Gemini rejected GEMINI_API_KEY. Check .env and rerun.")
                        conn.close()
                        return done
                    if not custom_scorer and ("404" in str(exc) or "NOT_FOUND" in str(exc)):
                        print(f"[score] model '{current_model}' unavailable (404); trying next fallback...")
                        rotate_model = True
                        break
                    if not custom_scorer and _is_daily_quota_error(exc):
                        print(f"[score] daily quota full on '{current_model}'; switching model...")
                        rotate_model = True
                        break
                    if _is_quota_error(exc) and attempt < 2:
                        parsed = _parse_retry_seconds(exc)
                        wait = 0 if sleep_sec == 0 else (min(max(parsed + 2, 15), 65) if parsed else 20 * (attempt + 1))
                        print(f"[score] rate limited on '{current_model}'; waiting {wait}s")
                        time.sleep(wait)
                        continue
                    if _is_quota_error(exc):
                        if not custom_scorer and len(active_models) > 1:
                            print(f"[score] '{current_model}' throttled/503 after retries; failing over...")
                            rotate_model = True
                        else:
                            exhausted_all = True
                        break
                    print(f"[score] error on job {job['id']}: {str(exc)[:180]}")
                    break

            if rotate_model:
                active_models.pop(0)
                if not active_models:
                    exhausted_all = True
                continue
            break

        if exhausted_all and fb_models and not fb_active:
            fb_active = True
            print(f"[score] all Gemini models out of quota; switching to free fallback provider ({FALLBACK_BASE_URL})")
            while fb_models and result is None:
                try:
                    api_calls += 1
                    result = score_via_openai_compat(job, profile, FALLBACK_BASE_URL, FALLBACK_API_KEY, fb_models[0])
                    used_model = f"fallback:{fb_models[0]}"
                except Exception as exc:
                    print(f"[score] fallback '{fb_models[0]}' failed ({str(exc)[:100]}); trying next fallback model")
                    fb_models.pop(0)
            if result is not None:
                exhausted_all = False
        if exhausted_all:
            print("[score] QUOTA_EXHAUSTED across all configured models; stopping cleanly. Unscored rows stay for next run.")
            conn.close()
            return done

        if result is None:
            continue

        # second opinion from a stronger model, only for promising rows (saves quota)
        if (not custom_scorer and rerank_on and not fb_active and result.match_score >= RERANK_FROM
                and RERANK_MODEL != used_model):
            try:
                api_calls += 1
                result = score_single_job(client, job, profile, model_name=RERANK_MODEL)
                used_model = RERANK_MODEL
            except Exception as exc:
                if _is_quota_error(exc):
                    rerank_on = False
                    print("[score] rerank model is out of quota; rerank OFF for the rest of this run")
                else:
                    print(f"[score] rerank skipped ({str(exc)[:80]}); keeping first score")

        cur.execute(
            "UPDATE jobs SET scored=1, match_score=?, priority=?, fit_reason=?, "
            "missing_skills=?, recommendation=?, resume_tip=?, scored_model=?, scored_at=? WHERE id=?",
            (result.match_score, result.priority, result.fit_reason,
             ", ".join(result.missing_skills), result.recommendation,
             getattr(result, "resume_tip", "") or "", used_model, now_utc(), job["id"]),
        )
        conn.commit()
        if sig:
            clone_cache[sig] = result
        done += 1
        print(f"  {result.match_score:>3}  {job['title'][:50]} @ {job['company'][:25]}")
        time.sleep(sleep_sec)

    conn.close()
    print(f"[score] scored {done} ({api_calls} API call(s))")
    return done


# ----------------------------------------------------------------------
# STAGE: EXCEL
# ----------------------------------------------------------------------
NOTES_HEADER = "Applied? (your notes)"
TIP_HEADER = "Resume / Pitch Tip"
EXCEL_HEADERS = [
    "ID", "Score", "Priority", "Recommendation", "Job Title", "Company", "Location",
    "Job Listing Link", "Direct Apply Link", "Company Link", "Fit Reason", "Missing Skills",
    "Source Channel", "Source Site", "From Email?", "Email Subject", "Email Sender",
    "Email Date", "In Digest?", "Digest Sent At", "Scored At", NOTES_HEADER, TIP_HEADER,
]
LINK_COLS = {8, 9, 10}


def _sheet(ws, rows: List[dict], pending_ids: set, pending_ts: str) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    hdr_fill = PatternFill("solid", start_color="1A73E8", end_color="1A73E8")
    green = PatternFill("solid", start_color="E6F4EA", end_color="E6F4EA")
    amber = PatternFill("solid", start_color="FEF7E0", end_color="FEF7E0")
    side = Side(style="thin", color="DADCE0")
    border = Border(left=side, right=side, top=side, bottom=side)
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    base_font = Font(name="Calibri", size=10)

    ws.append(EXCEL_HEADERS)
    ws.freeze_panes = "F2"
    ws.row_dimensions[1].height = 24
    for c in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = hdr_fill, Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_i, r in enumerate(rows, start=2):
        sent = r["emailed_in_digest"] == 1 or r["id"] in pending_ids
        sent_at = r["emailed_at"] or (pending_ts if r["id"] in pending_ids else "")
        ws.append([
            r["id"], r["match_score"] or 0, r["priority"] or "", r["recommendation"] or "",
            r["title"] or "", r["company"] or "", r["location"] or "",
            r["job_url"] or "", r["job_url_direct"] or r["job_url"] or "", r["company_url"] or "",
            r["fit_reason"] or "", r["missing_skills"] or "", r["source_channel"] or "",
            r["source_site"] or "", "Yes" if r["came_from_email"] else "No",
            r["source_email_subject"] or "", r["source_email_sender"] or "",
            r["source_email_date"] or "", "Yes" if sent else "No", sent_at, r["scored_at"] or "",
            r.get("user_notes") or "", r.get("resume_tip") or "",
        ])
        score = r["match_score"] or 0
        fill = green if score >= 80 else (amber if score >= MIN_SCORE else None)
        for c in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=r_i, column=c)
            cell.border = border
            if fill:
                cell.fill = fill
            url = str(cell.value or "")
            if c in LINK_COLS and url.startswith("http") and len(url) < 2000:
                cell.hyperlink = url
                cell.font = link_font
            else:
                cell.font = base_font
            cell.alignment = Alignment(vertical="top", wrap_text=c in (11, 12, 16, 23))
    widths = [6, 7, 10, 14, 34, 24, 20, 40, 40, 28, 50, 32, 18, 16, 10, 36, 28, 22, 10, 22, 22, 24, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions


def _applications_sheet(ws, app_rows: List[dict]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    headers = ["ID", "Status", "Applied Date", "Follow-Up Date", "Company", "Job Title", "Job Link", "Notes", "Updated At"]
    hdr_fill = PatternFill("solid", start_color="1A73E8", end_color="1A73E8")
    side = Side(style="thin", color="DADCE0")
    border = Border(left=side, right=side, top=side, bottom=side)
    link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    base_font = Font(name="Calibri", size=10)

    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = hdr_fill, Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_i, r in enumerate(app_rows, start=2):
        ws.append([
            r["id"], r["status"], r["applied_date"], r["follow_up_date"] or "",
            r["company"], r["title"], r["job_url"] or "", r["notes"] or "", r["updated_at"] or "",
        ])
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r_i, column=c)
            cell.border = border
            url = str(cell.value or "")
            if c == 7 and url.startswith("http") and len(url) < 2000:
                cell.hyperlink = url
                cell.font = link_font
            else:
                cell.font = base_font
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 8))

    for i, w in enumerate([6, 14, 14, 15, 24, 34, 40, 36, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions


def sync_notes_from_excel(db_path: str = DB_PATH, excel_path: str = EXCEL_PATH) -> int:
    """Sync user notes typed into Excel back into SQLite (and `applications`) before rebuilding."""
    import glob
    from openpyxl import load_workbook

    stem, ext = os.path.splitext(excel_path)
    files = [f for f in [excel_path] + glob.glob(f"{stem}_*{ext}") if os.path.isfile(f)]
    files.sort(key=os.path.getmtime)
    if not files:
        return 0
    conn = get_conn(db_path)
    current = {r["id"]: (r["user_notes"] or "") for r in conn.execute("SELECT id, user_notes FROM jobs")}
    changed = 0
    for path in files:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            print(f"[excel] could not read notes from '{path}': {exc}")
            continue
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header or "ID" not in header or NOTES_HEADER not in header:
                continue
            i_id, i_note = header.index("ID"), header.index(NOTES_HEADER)
            for row in rows:
                if len(row) <= max(i_id, i_note) or row[i_note] is None:
                    continue
                note = str(row[i_note]).strip()
                if not note or not isinstance(row[i_id], (int, float)):
                    continue
                jid = int(row[i_id])
                new = "" if note == "-" else note
                if jid in current and new != current[jid]:
                    current[jid] = new
                    conn.execute("UPDATE jobs SET user_notes=? WHERE id=?", (new, jid))
                    changed += 1
        wb.close()
    sync_applications_from_jobs(conn)
    conn.commit()
    conn.close()
    if changed:
        print(f"[excel] kept {changed} note(s) you typed in Excel (and synced to applications table)")
    return changed


def stage_excel(
    db_path: str = DB_PATH,
    excel_path: str = EXCEL_PATH,
    pending_ids: set = None,
    pending_ts: str = "",
) -> str:
    from openpyxl import Workbook

    pending_ids = pending_ids or set()
    sync_notes_from_excel(db_path, excel_path)
    conn = get_conn(db_path)
    sync_applications_from_jobs(conn)
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM jobs WHERE scored = 1 ORDER BY match_score DESC, id DESC"
    )]
    try:
        app_rows = [dict(r) for r in conn.execute(
            "SELECT * FROM applications ORDER BY applied_date DESC, id DESC"
        )]
    except sqlite3.Error:
        app_rows = []
    applied_keys = load_applied_keys(conn)
    conn.close()

    email_rows = [
        r for r in rows
        if r["came_from_email"] == 1 or r["emailed_in_digest"] == 1 or r["id"] in pending_ids
    ]
    # Top sheet: keep rows you marked yourself (so your tracker stays intact) but hide the copies of
    # an applied job that turned up on another site (they stay in All_Scored_Jobs).
    top_rows = [r for r in rows if (r["match_score"] or 0) >= MIN_SCORE
                and (note_is_applied(r.get("user_notes") or "")
                     or not is_applied_duplicate(r, applied_keys))]

    wb = Workbook()
    ws = wb.active
    ws.title = "All_Scored_Jobs"
    _sheet(ws, rows, pending_ids, pending_ts)
    _sheet(wb.create_sheet("From_Email_Or_Sent"), email_rows, pending_ids, pending_ts)
    _sheet(wb.create_sheet(f"Top_{MIN_SCORE}_Plus"), top_rows, pending_ids, pending_ts)
    _applications_sheet(wb.create_sheet("Applications_Tracker"), app_rows)

    try:
        wb.save(excel_path)
    except PermissionError:
        alt = excel_path.replace(".xlsx", f"_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
        wb.save(alt)
        print(f"[excel] '{excel_path}' is open in Excel; saved '{alt}' instead")
        excel_path = alt
    print(f"[excel] {excel_path}: {len(rows)} total, {len(email_rows)} email/sent, {len(top_rows)} top, {len(app_rows)} tracked applications")
    return excel_path


# ----------------------------------------------------------------------
# STAGE: DIGEST
# ----------------------------------------------------------------------
def build_digest_html(jobs: List[dict]) -> str:
    e = lambda v: html.escape(str(v or ""))  # noqa: E731
    rows = ""
    for j in jobs:
        color = "#137333" if j["match_score"] >= 80 else "#B06000"
        direct = j["job_url_direct"] or j["job_url"]
        origin = (
            f"Email alert: {e(j['source_email_subject'])}" if j["came_from_email"]
            else f"{e(j['source_channel'])}: {e(j['source_site'])}"
        )
        company = (
            f' | <a href="{e(j["company_url"])}">Company</a>'
            if str(j["company_url"] or "").startswith("http") else ""
        )
        posted = f" | Posted: {e(str(j.get('date_posted') or '')[:10])}" if j.get("date_posted") else ""
        tip_html = (
            f'<br><span style="color:#1A73E8"><b>Pitch tip:</b> {e(j.get("resume_tip"))}</span>'
            if j.get("resume_tip") else ""
        )
        rows += f"""<tr style="border-bottom:1px solid #DADCE0">
<td style="padding:10px;text-align:center;font-weight:bold;color:{color}">{j['match_score']}<br>
<span style="font-size:11px">{e(j['priority'])}</span></td>
<td style="padding:10px"><b>{e(j['title'])}</b><br>{e(j['company'])} - {e(j['location'])}<br>
<span style="font-size:11px;color:#575B5F">{origin}{posted}</span></td>
<td style="padding:10px;font-size:13px"><b>{e(j['recommendation'])}</b>: {e(j['fit_reason'])}<br>
<b>Missing:</b> {e(j['missing_skills']) or 'None'}{tip_html}</td>
<td style="padding:10px;white-space:nowrap"><a href="{e(j['job_url'])}">Listing</a> |
<a href="{e(direct)}">Apply</a>{company}</td></tr>"""
    return f"""<html><body style="font-family:Arial,sans-serif">
<h2 style="color:#1A73E8">{len(jobs)} new cloud job match(es), score &ge; {MIN_SCORE}</h2>
<p>Excel tracker attached. Add your application status in the <b>{e(NOTES_HEADER)}</b> column — your notes sync back automatically into the <b>applications</b> table on the next run.</p>
<table style="width:100%;border-collapse:collapse;border:1px solid #DADCE0">
<tr style="background:#1A73E8;color:#fff"><th>Score</th><th>Role</th><th>Assessment</th><th>Links</th></tr>
{rows}</table></body></html>"""


def stage_digest(db_path: str = DB_PATH, excel_path: str = EXCEL_PATH) -> int:
    conn = get_conn(db_path)
    jobs = [dict(r) for r in conn.execute(
        "SELECT * FROM jobs WHERE scored=1 AND match_score>=? AND emailed_in_digest=0 "
        "ORDER BY match_score DESC, id DESC",
        (MIN_SCORE,),
    )]
    applied_keys = load_applied_keys(conn)
    conn.close()

    # hide anything already applied to, including the same company+title found on another site
    before = len(jobs)
    jobs = [j for j in jobs if not is_applied_duplicate(j, applied_keys)]
    if before != len(jobs):
        print(f"[digest] hid {before - len(jobs)} job(s) already applied to (same company+title)")

    if not jobs:
        print(f"[digest] nothing new with score >= {MIN_SCORE}")
        stage_excel(db_path, excel_path)
        return 0

    ts = now_utc()
    ids = {j["id"] for j in jobs}
    attach_path = stage_excel(db_path, excel_path, pending_ids=ids, pending_ts=ts)

    if not (EMAIL_ADDRESS and EMAIL_PASSWORD):
        print("[digest] SMTP credentials missing; Excel written, email skipped.")
        return 0

    msg = MIMEMultipart()
    msg["From"], msg["To"] = EMAIL_ADDRESS, RECIPIENT_EMAIL
    msg["Subject"] = f"[Job Agent] {len(jobs)} new cloud match(es), score >= {MIN_SCORE}"
    msg.attach(MIMEText(build_digest_html(jobs), "html"))
    if os.path.exists(attach_path):
        with open(attach_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", f'attachment; filename="{os.path.basename(attach_path)}"')
        msg.attach(part)

    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=30) as server:
            server.ehlo()
            server.starttls()
            server.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
            server.sendmail(EMAIL_ADDRESS, [RECIPIENT_EMAIL], msg.as_string())
    except Exception as exc:
        print(f"[digest] SEND FAILED: {exc}. Rows left un-emailed; will retry next run.")
        stage_excel(db_path, excel_path)
        return 0

    conn = get_conn(db_path)
    conn.executemany(
        "UPDATE jobs SET emailed_in_digest=1, emailed_at=? WHERE id=?",
        [(ts, i) for i in ids],
    )
    conn.commit()
    conn.close()
    stage_excel(db_path, excel_path)
    print(f"[digest] sent {len(jobs)} job(s) to {RECIPIENT_EMAIL}")
    return len(jobs)


# ----------------------------------------------------------------------
# STAGE: STATUS (Terminal Dashboard)
# ----------------------------------------------------------------------
def stage_status(db_path: str = DB_PATH) -> None:
    conn = get_conn(db_path)
    total = conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
    scored = conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=1").fetchone()[0]
    unscored = conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=0").fetchone()[0]
    prefiltered = conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=1 AND match_score=0").fetchone()[0]
    dead_count = conn.execute(
        "SELECT COUNT(*) FROM jobs WHERE scored=1 AND scored_model='liveness-gate'"
    ).fetchone()[0]
    high = conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=1 AND match_score>=80").fetchone()[0]
    top = conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=1 AND match_score>=?", (MIN_SCORE,)).fetchone()[0]
    noted = conn.execute("SELECT COUNT(*) FROM jobs WHERE TRIM(COALESCE(user_notes, '')) != ''").fetchone()[0]
    try:
        app_total = conn.execute("SELECT COUNT(*) FROM applications").fetchone()[0]
    except sqlite3.Error:
        app_total = 0

    applied_keys = load_applied_keys(conn)
    top_unapplied = [
        r for r in conn.execute(
            "SELECT id, match_score, title, company, location, recommendation, user_notes "
            "FROM jobs WHERE scored=1 AND match_score>=? AND TRIM(COALESCE(user_notes, '')) = '' "
            "ORDER BY match_score DESC, id DESC",
            (MIN_SCORE,),
        ).fetchall()
        if not is_applied_duplicate(dict(r), applied_keys)
    ][:10]
    conn.close()

    print("\n=== CLOUD JOB AGENT STATUS ===")
    print(f"Database       : {db_path} ({total} total jobs)")
    print(f"Scored / Queue : {scored} scored ({prefiltered} pre-filtered, {dead_count} dead links) | {unscored} waiting")
    print(f"Matches        : {top} scoring >= {MIN_SCORE} ({high} scoring 80+)")
    print(f"Applications   : {app_total} in `applications` table ({noted} notes in `jobs`)")
    print(f"Model Chain    : {' -> '.join(MODEL_CHAIN)}")
    if top_unapplied:
        print(f"\nTop Unapplied Matches (score >= {MIN_SCORE}):")
        for r in top_unapplied:
            print(f"  [{r['match_score']:>3}] #{r['id']:<4} {r['title'][:40]} @ {r['company'][:22]} ({r['location'][:18]})")
    print("==============================\n")


# ----------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------
def stage_remote(db_path: str = DB_PATH) -> int:
    """Worldwide remote jobs from public feeds (see remote_feeds.py). Skipped if that file is absent."""
    try:
        import remote_feeds
    except ImportError:
        print("[remote] remote_feeds.py not found; skipping.")
        return 0
    refresh_applied_cache(db_path)
    return remote_feeds.run(db_path, get_conn, insert_job, job_hash, SENIOR_TITLE)


STAGES = {
    "init": stage_init, "scrape": stage_scrape, "ats": stage_ats, "remote": stage_remote, "email": stage_email,
    "score": stage_score, "excel": stage_excel, "digest": stage_digest, "status": stage_status,
}
ORDER = ["init", "scrape", "ats", "remote", "email", "score", "digest"]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--stage", choices=sorted(STAGES), help="run a single stage")
    ap.add_argument("--max", type=int, default=None, help="override MAX_SCORE_PER_RUN for scoring")
    ap.add_argument(
        "--track",
        type=str,
        default=None,
        help="Track an application directly: 'Company | Title | applied | Optional note'",
    )
    args = ap.parse_args()

    if args.track:
        parts = [p.strip() for p in args.track.split("|")]
        if len(parts) < 2:
            raise SystemExit("Usage: python job_agent.py --track \"Company | Job Title | applied | Note\"")
        comp, ttl = parts[0], parts[1]
        st = parts[2] if len(parts) >= 3 and parts[2] else "applied"
        nts = parts[3] if len(parts) >= 4 else f"Applied {today_iso()}"
        key = upsert_application(DB_PATH, comp, ttl, status=st, notes=nts)
        print(f"[track] recorded [{st.upper()}] {comp} - {ttl} (key: {key})")
        return

    if args.stage:
        if args.stage != "init":
            stage_init()
        if args.stage == "score" and args.max is not None:
            stage_score(max_calls=args.max)
        else:
            STAGES[args.stage]()
        return
    for name in ORDER:
        try:
            if name == "score" and args.max is not None:
                stage_score(max_calls=args.max)
            else:
                STAGES[name]()
        except SystemExit:
            raise
        except Exception as exc:
            print(f"[{name}] FAILED: {exc}", file=sys.stderr)


if __name__ == "__main__":
    main()