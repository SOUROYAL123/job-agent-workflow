"""
followup_reminder.py - emails YOU a list of follow-ups that are due, with each message ready to copy.

It sends nothing to recruiters. Two sources:
  1. networking_kit.xlsx: rows with Status = Sent and a Date sent 6+ days ago.
  2. Your jobs file notes: rows whose note says "Applied <date>" (e.g. "Applied 25 Oct") 6+ days ago,
     unless the note says followed up, interview, replied, offer, rejected or no reply.
After you send a follow-up, change the Status to "Followed up" (or add "followed up" to the note).
At most one reminder email per day. Run: python followup_reminder.py   (add --print to preview only)
"""
import os
import re
import smtplib
import sqlite3
import sys
from datetime import date, datetime
from email.mime.text import MIMEText

from openpyxl import load_workbook

try:
    from dotenv import load_dotenv
    load_dotenv(override=True)
except Exception:
    pass

from networking_helper import follow_up, parse_sent_date  # noqa: E402

DB_PATH = os.getenv("DB_PATH", "jobs.db")
KIT_PATH = os.getenv("NETWORKING_KIT_PATH", "networking_kit.xlsx")
GUARD = "followup_last_sent.txt"
DUE_DAYS = 6
EMAIL_ADDRESS = os.getenv("EMAIL_ADDRESS", "").strip()
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "").strip()
RECIPIENT = os.getenv("RECIPIENT_EMAIL", "").strip() or EMAIL_ADDRESS
SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))

MONTHS = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
DM = re.compile(rf"(\d{{1,2}})(?:st|nd|rd|th)?\s*({MONTHS})[a-z]*", re.I)
MD = re.compile(rf"({MONTHS})[a-z]*\s*(\d{{1,2}})", re.I)


def date_from_note(note: str):
    m = DM.search(note)
    day, mon = (m.group(1), m.group(2)) if m else (None, None)
    if not m:
        m = MD.search(note)
        if not m:
            return None
        mon, day = m.group(1), m.group(2)
    return parse_sent_date(f"{int(day)} {mon[:3].title()}")


def due_from_kit(today):
    items = []
    if not os.path.exists(KIT_PATH):
        return items
    try:
        wb = load_workbook(KIT_PATH, data_only=True)
    except Exception as exc:
        print(f"[followup] cannot read {KIT_PATH}: {exc}")
        return items
    ws = wb["Outreach_Drafts"]
    hdr = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[ix["Status"]] or "").strip().lower() != "sent":
            continue
        sent = parse_sent_date(row[ix["Date sent"]])
        if not sent or (today - sent).days < DUE_DAYS:
            continue
        who = row[ix["Contact name"]] or "[Name]"
        msg = follow_up(row[ix["Job title"]], row[ix["Company"]]).replace("[Name]", str(who)).replace("[date]", sent.strftime("%d %b"))
        items.append((f"{who} at {row[ix['Company']]} - {row[ix['Job title']]}", (today - sent).days, msg, row[ix["Job link"]]))
    return items


def due_from_applications(today):
    items = []
    if not os.path.exists(DB_PATH):
        return items
    conn = sqlite3.connect(DB_PATH)
    for title, company, url, note in conn.execute(
            "SELECT title, company, COALESCE(NULLIF(job_url_direct, ''), job_url), user_notes FROM jobs "
            "WHERE user_notes LIKE '%pplied%'"):
        note = note or ""
        if re.search(r"followed|interview|reject|declin|replied|offer|not selected|no reply|no response|withdr", note, re.I):
            continue  # already followed up, or the company has already answered
        d = date_from_note(note)
        if not d or (today - d).days < DUE_DAYS:
            continue
        msg = follow_up(title, company).replace("[date]", d.strftime("%d %b"))
        items.append((f"Application: {title} at {company}", (today - d).days, msg, url))
    conn.close()
    return items


def build_body(items):
    lines = [f"{len(items)} follow-up(s) are due. Send each one yourself, then update the Status/notes.\n"]
    for i, (label, days, msg, link) in enumerate(items, 1):
        lines += [f"{i}. {label}  ({days} days since you applied or messaged)", f"   Link: {link}", "   Message to copy:", f"   {msg}", ""]
    lines.append("Replace any [Name] before sending. Send one follow-up only; if there is no reply, move on.")
    return "\n".join(lines)


def main():
    preview = "--print" in sys.argv
    force = "--force" in sys.argv
    today = date.today()
    if not preview and not force and os.path.exists(GUARD) and open(GUARD).read().strip() == today.isoformat():
        print("[followup] already sent today; skipping.")
        return 0
    kit_items = due_from_kit(today)
    seen = {i[3] for i in kit_items}
    items = kit_items + [i for i in due_from_applications(today) if i[3] not in seen]
    if not items:
        print("[followup] nothing due today.")
        return 0
    body = build_body(items)
    if preview:
        print(body)
        return len(items)
    if not (EMAIL_ADDRESS and EMAIL_PASSWORD):
        print("[followup] email settings missing; printing instead.\n" + body)
        return len(items)
    msg = MIMEText(body, "plain", "utf-8")
    msg["From"], msg["To"] = EMAIL_ADDRESS, RECIPIENT
    msg["Subject"] = f"[Job Agent] {len(items)} follow-up(s) due today"
    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=30) as s:
            s.ehlo(); s.starttls(); s.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
            s.sendmail(EMAIL_ADDRESS, [RECIPIENT], msg.as_string())
        open(GUARD, "w").write(today.isoformat())
        print(f"[followup] reminder emailed: {len(items)} due.")
    except Exception as exc:
        print(f"[followup] email failed: {str(exc)[:100]}")
    return len(items)


if __name__ == "__main__":
    main()
