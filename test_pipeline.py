"""Offline tests: no network, no Gmail, no Gemini. Run: python test_pipeline.py"""
import os
import sys
import tempfile
from email.message import EmailMessage

tmp = tempfile.mkdtemp()
os.environ.update(DB_PATH=os.path.join(tmp, "t.db"), EXCEL_BACKUP_PATH=os.path.join(tmp, "t.xlsx"),
                  EMAIL_ADDRESS="me@example.com", EMAIL_PASSWORD="x", RECIPIENT_EMAIL="me@example.com",
                  GEMINI_API_KEY="")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import job_agent as ja  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

# Use the temp paths directly: job_agent calls load_dotenv(override=True), which would otherwise
# swap os.environ["DB_PATH"] for the REAL jobs.db from .env and let this test write into it.
DB, XL = os.path.join(tmp, "t.db"), os.path.join(tmp, "t.xlsx")
ja.EMAIL_ADDRESS, ja.EMAIL_PASSWORD, ja.RECIPIENT_EMAIL = "me@example.com", "x", "me@example.com"
ok = lambda name: print(f"PASS  {name}")  # noqa: E731


# 1. Hashing: the original Indeed bug
h1 = ja.job_hash("https://in.indeed.com/viewjob?jk=aaa111")
h2 = ja.job_hash("https://in.indeed.com/viewjob?jk=bbb222")
assert h1 != h2
assert ja.job_hash("https://in.indeed.com/viewjob?jk=aaa111&utm=x") == h1
assert ja.job_hash("https://www.linkedin.com/jobs/view/cloud-engineer-3912345678?trk=a") == \
       ja.job_hash("https://in.linkedin.com/jobs/view/3912345678?refId=zz")
assert ja.job_hash("https://www.naukri.com/job-listings-aws-111") != ja.job_hash("https://www.naukri.com/job-listings-aws-222")
ok("hashing: Indeed ids distinct, LinkedIn tracking params ignored")

# 2. Email parsing
def make(sender, subject, html):
    m = EmailMessage()
    m["From"], m["Subject"], m["Date"] = sender, subject, "Sat, 24 Oct 2026 08:00:00 +0530"
    m.set_content("plain")
    m.add_alternative(html, subtype="html")
    return m

alert = make("LinkedIn Job Alerts <jobalerts@linkedin.com>", "5 new jobs for AWS Cloud Support",
             '<a href="https://www.linkedin.com/comm/jobs/view/3900000001?trk=x">Cloud Support Engineer</a>'
             '<a href="https://www.linkedin.com/comm/jobs/view/3900000001?trk=y">Cloud Support Engineer</a>'
             '<a href="https://in.indeed.com/rc/clk?jk=abc111">AWS Support Associate</a>'
             '<a href="https://in.indeed.com/rc/clk?jk=abc222">Implementation Specialist</a>'
             '<a href="https://linkedin.com/unsubscribe">Unsubscribe</a>')
found = ja.parse_alert_message(alert)
assert len(found) == 3, len(found)  # duplicate collapsed, unsubscribe dropped, both Indeed ids kept
personal = make("Friend <friend@gmail.com>", "dinner saturday?", '<a href="https://example.com">menu</a>')
assert ja.parse_alert_message(personal) == []
ok("email parsing: 3 unique jobs, personal mail ignored")

# 3. IMAP uses PEEK and only flags alert mails
class FakeIMAP:
    fetch_args, stored = [], []
    def __init__(self, *a): pass
    def login(self, *a): pass
    def select(self, *a): pass
    def search(self, *a): return "OK", [b"1 2"]
    def fetch(self, e_id, spec):
        FakeIMAP.fetch_args.append(spec)
        return "OK", [(b"hdr", (alert if e_id == b"1" else personal).as_bytes())]
    def store(self, e_id, *a): FakeIMAP.stored.append(e_id)
    def logout(self): pass

ja.stage_init(DB)
ja.imaplib.IMAP4_SSL = FakeIMAP
added = ja.stage_email(DB)
assert added == 3, added
assert all("PEEK" in s for s in FakeIMAP.fetch_args), FakeIMAP.fetch_args
assert FakeIMAP.stored == [b"1"], FakeIMAP.stored
ok("IMAP: BODY.PEEK used; only the alert email marked Seen")

# 4. Scoring: fake scorer, senior skip, quota stop
import sqlite3  # noqa: E402
conn = sqlite3.connect(DB)
for i, t in enumerate(["Senior DevOps Engineer", "L1 Cloud Support Engineer", "AWS Implementation Specialist"]):
    ja.insert_job(conn.cursor(), job_hash=ja.job_hash(f"https://x.com/{i}"), title=t, company=f"Co{i}",
                  location="Kolkata", description="d", job_url=f"https://x.com/{i}",
                  job_url_direct=f"https://x.com/{i}", company_url="https://x.com",
                  source_channel="JobSpy Scraper", source_site="indeed")
conn.commit(); conn.close()

calls = []
def fake_scorer(job):
    calls.append(job["title"])
    s = 85 if "Support" in job["title"] else 65 if "Implementation" in job["title"] else 30
    return ja.JobMatch(match_score=s, priority="High" if s >= 80 else "Medium", fit_reason="ok & <b>fine</b>",
                       missing_skills=["Linux"], recommendation="Apply now")
n = ja.stage_score(DB, scorer=fake_scorer, sleep_sec=0)
conn = sqlite3.connect(DB)
assert conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=0").fetchone()[0] == 0
assert "Senior DevOps Engineer" not in calls  # skipped without an API call
assert conn.execute("SELECT match_score FROM jobs WHERE title LIKE 'Senior%'").fetchone()[0] == 0
conn.close()
ok(f"scoring: {n} rows scored, senior title skipped without API call")

# quota: stops cleanly and leaves rows unscored
conn = sqlite3.connect(DB)
ja.insert_job(conn.cursor(), job_hash="q1", title="Cloud Ops", company="Q", job_url="https://q.com/1",
              source_channel="t", description="d")
conn.commit(); conn.close()
def quota(job): raise RuntimeError("429 RESOURCE_EXHAUSTED")
ja.time.sleep = lambda s: None
ja.stage_score(DB, scorer=quota, sleep_sec=0)
conn = sqlite3.connect(DB)
assert conn.execute("SELECT scored FROM jobs WHERE job_hash='q1'").fetchone()[0] == 0
conn.close()
ok("scoring: quota error stops run, row stays unscored for retry")

# 5. Digest: SMTP failure leaves rows un-emailed; success marks them and attachment shows Sent
class BadSMTP:
    def __init__(self, *a, **k): pass
    def __enter__(self): return self
    def __exit__(self, *a): pass
    def ehlo(self): pass
    def starttls(self): pass
    def login(self, *a): raise RuntimeError("auth blocked")
ja.smtplib.SMTP = BadSMTP
assert ja.stage_digest(DB, XL) == 0
conn = sqlite3.connect(DB)
assert conn.execute("SELECT COUNT(*) FROM jobs WHERE emailed_in_digest=1").fetchone()[0] == 0
conn.close()
ok("digest: SMTP failure does not crash and does not mark rows sent")

captured = {}
class GoodSMTP(BadSMTP):
    def login(self, *a): pass
    def sendmail(self, frm, to, body): captured["body"] = body
ja.smtplib.SMTP = GoodSMTP
sent = ja.stage_digest(DB, XL)
assert sent == 5, sent  # 2 scraped + 3 email-alert jobs all score >= 60 with the fake scorer
assert "&lt;b&gt;fine" in __import__("email").message_from_string(captured["body"]).get_payload()[0].get_payload(decode=True).decode()
conn = sqlite3.connect(DB)
assert conn.execute("SELECT COUNT(*) FROM jobs WHERE emailed_in_digest=1").fetchone()[0] == 5
conn.close()
assert ja.stage_digest(DB, XL) == 0  # nothing re-sent
ok("digest: sent once, HTML-escaped, never re-sent")

# 6. Excel: hyperlinks, sheets, sent flags
wb = load_workbook(XL)
assert wb.sheetnames == ["All_Scored_Jobs", "From_Email_Or_Sent", "Top_60_Plus"], wb.sheetnames
ws = wb["All_Scored_Jobs"]
hdr = [c.value for c in ws[1]]
link_cell = ws.cell(row=2, column=hdr.index("Job Listing Link") + 1)
assert link_cell.hyperlink is not None and link_cell.hyperlink.target.startswith("http")
in_digest = [ws.cell(row=r, column=hdr.index("In Digest?") + 1).value for r in range(2, ws.max_row + 1)]
assert in_digest.count("Yes") == 5 and in_digest.count("No") == 1  # senior-skipped row not sent
assert len(wb["Top_60_Plus"]["A"]) - 1 == 5
assert wb["From_Email_Or_Sent"].max_row - 1 == 5
ok("excel: 3 sheets, clickable links, Sent flags correct")


# 7. Title filter: architect fix
assert not ja.SENIOR_TITLE.search("Associate Solutions Architect")
assert not ja.SENIOR_TITLE.search("Junior Cloud Architect")
assert ja.SENIOR_TITLE.search("Enterprise Architect")
assert ja.SENIOR_TITLE.search("Senior Cloud Support Engineer")
print("PASS  title filter: junior/associate architect kept, senior/enterprise skipped")

# 8. Notes survive the rebuild (the "Excel overwrite trap")
from openpyxl import load_workbook as _lw  # noqa: E402
wb = _lw(XL)
ws = wb["Top_60_Plus"]
hdr = [c.value for c in ws[1]]
ni, ii = hdr.index(ja.NOTES_HEADER) + 1, hdr.index("ID") + 1
top_ids = [ws.cell(row=r, column=ii).value for r in range(2, ws.max_row + 1)]
a, b = top_ids[0], top_ids[1]
ws.cell(row=2, column=ni).value = "Applied 24 Oct, Sprinto portal"
wb["All_Scored_Jobs"].cell(row=1 + 1 + 0, column=ni)  # untouched sheet stays blank for this row
wb.save(XL)
ja.stage_excel(DB, XL)
wb2 = _lw(XL)
for name in ("All_Scored_Jobs", "From_Email_Or_Sent", "Top_60_Plus"):
    w = wb2[name]
    h = [c.value for c in w[1]]
    got = {w.cell(row=r, column=h.index("ID") + 1).value: w.cell(row=r, column=h.index(ja.NOTES_HEADER) + 1).value
           for r in range(2, w.max_row + 1)}
    if a in got:
        assert got[a] == "Applied 24 Oct, Sprinto portal", (name, got[a])
conn = sqlite3.connect(DB)
assert conn.execute("SELECT user_notes FROM jobs WHERE id=?", (a,)).fetchone()[0] == "Applied 24 Oct, Sprinto portal"
conn.close()
ok("notes: typed in one sheet, kept in DB and shown on every sheet after rebuild")

# clear with "-"
wb = _lw(XL); ws = wb["Top_60_Plus"]
ws.cell(row=2, column=ni).value = "-"
wb.save(XL)
ja.stage_excel(DB, XL)
conn = sqlite3.connect(DB)
assert conn.execute("SELECT user_notes FROM jobs WHERE id=?", (a,)).fetchone()[0] == ""
conn.close()
ok("notes: '-' clears a note")

# old database (no user_notes column) is migrated
old = os.path.join(tmp, "old.db")
ja.stage_init(old)                      # same schema the first zip created...
c = sqlite3.connect(old)
c.execute("ALTER TABLE jobs DROP COLUMN user_notes")   # ...minus the new column
c.commit()
assert "user_notes" not in [r[1] for r in c.execute("PRAGMA table_info(jobs)")]
c.close()
ja.stage_init(old)
c = sqlite3.connect(old)
assert "user_notes" in [r[1] for r in c.execute("PRAGMA table_info(jobs)")]
c.close()
ok("migration: old database gains user_notes column")


# 9. Bad API key stops after ONE call instead of failing every job
conn = sqlite3.connect(DB)
for i in range(3):
    ja.insert_job(conn.cursor(), job_hash=f"auth{i}", title=f"Cloud Support {i}", company="A",
                  job_url=f"https://a.com/{i}", source_channel="t", description="d")
conn.commit(); conn.close()
auth_calls = []
def bad_key(job):
    auth_calls.append(1)
    raise RuntimeError("400 INVALID_ARGUMENT. API key not valid. API_KEY_INVALID")
ja.stage_score(DB, scorer=bad_key, sleep_sec=0)
assert len(auth_calls) == 1, len(auth_calls)
conn = sqlite3.connect(DB)
assert conn.execute("SELECT COUNT(*) FROM jobs WHERE scored=0 AND job_hash LIKE 'auth%'").fetchone()[0] == 3
conn.close()
ok("auth: invalid key stops after 1 call, all rows stay unscored")

# 10. .env beats a stale environment variable
import shutil, subprocess, textwrap  # noqa: E402
d = tempfile.mkdtemp()
shutil.copy(os.path.join(os.path.dirname(os.path.abspath(__file__)), "job_agent.py"), d)
open(os.path.join(d, ".env"), "w").write("GEMINI_API_KEY=from_dotenv_file\n")
env = dict(os.environ, GEMINI_API_KEY="stale_system_value")
out = subprocess.run([sys.executable, "-c", "import job_agent as j; print(j.GEMINI_API_KEY)"],
                     cwd=d, env=env, capture_output=True, text=True).stdout.strip()
assert out == "from_dotenv_file", out
ok(".env: overrides a stale GEMINI_API_KEY environment variable")

print("\nALL TESTS PASSED")
