"""
networking_helper.py - builds your referral / recruiter outreach kit from the jobs the agent scored.

It does NOT message anyone and does NOT touch LinkedIn. It writes networking_kit.xlsx with:
  Outreach_Drafts   - one row per strong job: company, link, LinkedIn people-search links you click
                      yourself, a connection note (<=300 chars), a referral ask, a follow-up date.
  Target_Companies  - one row per company, with search links for recruiters and employees.
  How_To_Use        - the sending rules.
Your own columns (Contact, Status, Date sent, Notes) are kept when the file is rebuilt.
Run: python networking_helper.py
"""
import os
import re
import sqlite3
import sys
from datetime import date, timedelta
from urllib.parse import quote_plus

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv
    load_dotenv(override=True)
except Exception:
    pass

DB_PATH = os.getenv("DB_PATH", "jobs.db")
OUT_PATH = os.getenv("NETWORKING_KIT_PATH", "networking_kit.xlsx")
MIN_SCORE = int(os.getenv("NETWORK_MIN_SCORE", os.getenv("MIN_EMAIL_MATCH_SCORE", "60")))
FOLLOW_UP_DAYS = 6

# Employers named in your profile; always included even if no job is open right now.
ALWAYS_TARGET = ["Cloud4C", "Minfy", "CloudThat"]

USER_COLS = ["Contact name", "Status", "Date sent", "Notes"]
STATUS_HELP = "Not sent / Sent / Followed up / Replied / Referred / No reply"


def parse_sent_date(value):
    """Accept an Excel date, an ISO date, or text like '26 Oct', '26-10-2026', 'Oct 26'. Return date or None."""
    from datetime import datetime
    if value is None or value == "":
        return None
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    for fmt in ("%d %b", "%d %B", "%b %d", "%B %d"):
        try:
            d = datetime.strptime(text, fmt).date().replace(year=date.today().year)
            return d if d <= date.today() else d.replace(year=d.year - 1)
        except ValueError:
            pass
    return None


RESUME_SUPPORT = "Sourav_Roy_Chowdhury_AWS_Cloud_Support.pdf"
RESUME_IMPL = "Sourav_Roy_Chowdhury_AWS_Cloud_Implementation.pdf"
IMPL_WORDS = re.compile(r"implementation|onboarding|solutions? (engineer|consultant|architect)|integration|"
                        r"deployment|customer success|consultant|professional services|migration", re.I)
SUPPORT_WORDS = re.compile(r"support|noc\b|operations|incident|help ?desk|troubleshoot|ticket|monitoring|"
                           r"sre\b|reliability|l1|l2|service desk", re.I)


def pick_resume(title: str, description: str = "") -> str:
    """Title decides first; if the title is neutral, the first 1500 characters of the description vote."""
    t = title or ""
    if IMPL_WORDS.search(t) and not SUPPORT_WORDS.search(t):
        return RESUME_IMPL
    if SUPPORT_WORDS.search(t) and not IMPL_WORDS.search(t):
        return RESUME_SUPPORT
    d = (description or "")[:1500]
    impl, sup = len(IMPL_WORDS.findall(t + " " + d)), len(SUPPORT_WORDS.findall(t + " " + d))
    return RESUME_IMPL if impl > sup else RESUME_SUPPORT


def li_people(query: str) -> str:
    return "https://www.linkedin.com/search/results/people/?keywords=" + quote_plus(query)


def clean_company(name: str) -> str:
    name = (name or "").strip()
    return "" if (not name or name.lower().startswith("via ")) else name


def short(text: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", (text or "").strip())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def connection_note(title: str, company: str) -> str:
    note = (
        "Hi [Name], I'm a Kolkata-based career changer moving into AWS cloud support "
        "(Terraform two-AZ project, Linux and network troubleshooting). "
        f"I'm interested in the {short(title, 45)} role at {short(company, 30)}. "
        "Happy to connect."
    )
    return note if len(note) <= 300 else short(note, 300)


def referral_ask(title: str, company: str, url: str) -> str:
    return (
        f"Hi [Name], thank you for connecting. I'm applying for the {title} role at {company} ({url}). "
        "My background is an AWS and Terraform project, Linux and networking troubleshooting, and "
        "14 years of client-facing production delivery. I haven't held a paid cloud role yet, so I'd value "
        "your honest view. If you think I'm a fit, would you be open to referring me or pointing me to "
        "the right recruiter? I can send my resume and GitHub link. Thank you for your time."
    )


def follow_up(title: str, company: str) -> str:
    return (
        f"Hi [Name], I applied for the {title} role at {company} on [date]. My background combines an AWS "
        "and Terraform project, Linux and networking troubleshooting, and 14 years of client-facing "
        "production delivery. Could you tell me where the shortlisting stands? Thank you for your time."
    )


def load_jobs():
    if not os.path.exists(DB_PATH):
        sys.exit(f"Database not found: {DB_PATH}. Run the job agent first.")
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT id, title, company, location, match_score, job_url, job_url_direct, user_notes, description "
        "FROM jobs WHERE scored=1 AND match_score>=? ORDER BY match_score DESC, id DESC",
        (MIN_SCORE,),
    ).fetchall()
    conn.close()
    return rows


def read_user_cols(path: str, sheet: str, key_col: str) -> dict:
    """Return {key: {col: value}} for the columns you typed, so a rebuild never loses them."""
    keep = {}
    if not os.path.exists(path):
        return keep
    try:
        ws = load_workbook(path)[sheet]
    except Exception:
        return keep
    hdr = [c.value for c in ws[1]]
    if key_col not in hdr:
        return keep
    ki = hdr.index(key_col)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[ki] is None:
            continue
        keep[str(row[ki])] = {c: row[hdr.index(c)] for c in USER_COLS if c in hdr}
    return keep


def style(ws, widths):
    head = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main():
    jobs = load_jobs()
    saved_jobs = read_user_cols(OUT_PATH, "Outreach_Drafts", "Job ID")
    saved_cos = read_user_cols(OUT_PATH, "Target_Companies", "Company")
    today = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Outreach_Drafts"
    ws.append(["Job ID", "Score", "Company", "Job title", "Location", "Resume to send", "Job link",
               "Find recruiters (click)", "Find employees (click)", "Connection note (<=300 chars)",
               "Referral ask (after they accept)", "Follow-up message", "Follow up on"] + USER_COLS)
    companies = {}
    for r in jobs:
        co = clean_company(r["company"])
        if not co:
            continue
        link = r["job_url_direct"] or r["job_url"]
        companies.setdefault(co, []).append(r["title"])
        prev = saved_jobs.get(str(r["id"]), {})
        sent = parse_sent_date(prev.get("Date sent"))
        due = (sent + timedelta(days=FOLLOW_UP_DAYS)) if sent else (today + timedelta(days=FOLLOW_UP_DAYS))
        ws.append([r["id"], r["match_score"], co, r["title"], r["location"],
                   pick_resume(r["title"], r["description"]), link,
                   li_people(f"{co} recruiter OR talent acquisition"),
                   li_people(f"{co} cloud support engineer"),
                   connection_note(r["title"], co), referral_ask(r["title"], co, link),
                   follow_up(r["title"], co), due.isoformat(),
                   prev.get("Contact name") or "", prev.get("Status") or "Not sent",
                   prev.get("Date sent") or "", prev.get("Notes") or ""])
    style(ws, [8, 7, 20, 30, 16, 34, 40, 28, 28, 55, 70, 60, 12, 18, 14, 12, 30])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
        for c in row:
            if c.value and str(c.value).startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")

    wc = wb.create_sheet("Target_Companies")
    wc.append(["Company", "Open strong roles", "Find recruiters (click)", "Find cloud support staff (click)",
               "Find alumni/Kolkata contacts (click)"] + USER_COLS)
    for co in sorted(set(companies) | set(ALWAYS_TARGET), key=str.lower):
        prev = saved_cos.get(co, {})
        wc.append([co, len(companies.get(co, [])),
                   li_people(f"{co} recruiter OR talent acquisition"),
                   li_people(f"{co} cloud support"),
                   li_people(f"{co} Kolkata"),
                   prev.get("Contact name") or "", prev.get("Status") or "Not sent",
                   prev.get("Date sent") or "", prev.get("Notes") or ""])
    style(wc, [24, 12, 40, 40, 40, 18, 14, 12, 30])
    for row in wc.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            if c.value and str(c.value).startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")

    wh = wb.create_sheet("How_To_Use")
    for line in [
        ["How to use this kit"],
        ["1. Open a link in 'Find recruiters', pick a real person who recruits for cloud or IT roles at that company."],
        ["2. Replace [Name] in the connection note, send it yourself on LinkedIn. Aim for 5 to 10 people a day, not hundreds."],
        ["3. After they accept, send the referral ask. Add the job link only if you have applied or are about to."],
        ["4. Write their name, the Status and Date sent. 'Follow up on' recalculates 6 days after Date sent."],
        ["5. Send the follow-up once. If there is no reply, move on. Never message the same person repeatedly."],
        ["6. Keep every claim true: planned for November 2026 exam, personal AWS project, no paid cloud role yet."],
        ["Status values: " + STATUS_HELP],
        ["This file is rebuilt on every run. Your Contact, Status, Date sent and Notes are kept."],
    ]:
        wh.append(line)
    wh["A1"].font = Font(bold=True, size=13)
    wh.column_dimensions["A"].width = 120

    try:
        wb.save(OUT_PATH)
    except PermissionError:
        sys.exit(f"Close {OUT_PATH} in Excel and run again.")
    print(f"[network] {OUT_PATH}: {ws.max_row - 1} job(s), {wc.max_row - 1} company(ies)")


if __name__ == "__main__":
    main()
